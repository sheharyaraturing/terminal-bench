"""Fraction of the no-contact cohort identified.

Scores partial work: 18 of the 21 customers earns 0.857, not 0. Rows the agent
invented, and rows it repeated, count against the denominator, so neither padding
the sheet nor failing to collapse a repeated customer_id can inflate the score.
An agent that produces nothing reads 0.
"""

from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/no_contact.xlsx")
COLUMNS = ("customer_id", "customer_name", "segment", "data_issue")

# Sheets already named in a formula diagnostic: one line per affected sheet, not
# one per cell and not one per criterion.
_FLAGGED: set = set()


def _norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _sheet(path: Path):
    """The answer sheet as (header index map, data rows).

    Scans every sheet rather than trusting .active: an agent that leaves an empty
    Sheet1 in front of its answer is still correct. Formula cells are read back
    from a second, non-cached load so they surface as formulas instead of blanks.
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    if not path.exists():
        return None, []
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"cohort: could not open {path}: {type(exc).__name__}: {exc}")
        return None, []
    try:
        raw = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        raw = cached

    for title in sorted(cached.sheetnames, key=lambda t: (t != "NoContact", t)):
        rows = [list(r) for r in cached[title].iter_rows(values_only=True)]
        if title in raw.sheetnames:
            other = [list(r) for r in raw[title].iter_rows(values_only=True)]
            for i, row in enumerate(rows):
                if i >= len(other):
                    break
                for j, value in enumerate(row):
                    if value is None and j < len(other[i]):
                        rows[i][j] = other[i][j]
                        text = _norm(other[i][j])
                        if text.startswith("=") and (path, title) not in _FLAGGED:
                            _FLAGGED.add((path, title))
                            coord = f"{get_column_letter(j + 1)}{i + 1}"
                            print(f"cohort: {title}!{coord} holds the formula "
                                  f"{text!r} and no cached value; section 8 of "
                                  "metrics/Definitions.md requires recorded values, "
                                  "not formulas")
        for start, row in enumerate(rows[:5]):
            header = [_norm(c).lower() for c in row]
            if all(c in header for c in COLUMNS):
                idx = {c: header.index(c) for c in COLUMNS}
                body = [r for r in rows[start + 1:] if r and any(c is not None for c in r)]
                return idx, body
    print(f"cohort: no sheet in {path.name} carried all of {COLUMNS}")
    return None, []


def _ids(path: Path) -> tuple[list[str], int]:
    idx, body = _sheet(path)
    if idx is None:
        return [], 0
    ids = []
    for row in body:
        cell = _norm(row[idx["customer_id"]]) if idx["customer_id"] < len(row) else ""
        if cell:
            ids.append(cell)
    return ids, len(body)


@criterion(description="customers without support contact in the reporting period")
def cohort(workspace: Path) -> float:
    expected, _ = _ids(EXPECTED)
    if not expected:
        print("cohort: the answer key could not be read")
        return 0.0
    got, rows = _ids(workspace / "no_contact.xlsx")
    got_set = set(got)
    hits = sum(1 for cid in set(expected) if cid in got_set)
    return hits / max(len(set(expected)), rows, len(got))
