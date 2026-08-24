"""The two deliverables exist and carry the shape the specification asks for.

Shape only — whether the figures are right is graded by the cohort, attributes and
workings dimensions. Shape here means what section 8 and section 9 of
`metrics/Definitions.md` actually state: the sheet is named `NoContact`, the four
columns appear in the stated order, the sheet carries at least one data row, the
rows are sorted ascending by `customer_id`, and every one of the ten workings keys
is present and filled in. A headers-only shell satisfies none of them, so an empty
scaffold banks nothing.
"""

import json
from pathlib import Path

from rewardkit import criterion

SHEET = "NoContact"
SHEET_COLUMNS = ("customer_id", "customer_name", "segment", "data_issue")
WORKINGS_STRINGS = ("metric", "period_start", "period_end")
WORKINGS_NUMBERS = (
    "customers_total", "customers_with_contact", "value", "contact_rate",
    "orphan_tickets", "soft_deleted_tickets", "out_of_period_tickets",
)
WORKINGS_KEYS = WORKINGS_STRINGS + WORKINGS_NUMBERS

# Sheets already named in a formula diagnostic. One line per affected sheet is a
# useful pointer; one per cell, or one per criterion, is a flooded log.
_FLAGGED: set = set()


def _norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _grids(path: Path):
    """Every sheet, as a list of row lists, with formula cells filled back in.

    A workbook written by a program that never ran Excel has no cached formula
    values, so data_only=True reads those cells as None. Load a second time
    without it and merge, so a formula-written answer surfaces as its formula
    text and is named as such rather than being reported as an empty cell.
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    try:
        cached = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"artifacts: could not open {path}: {type(exc).__name__}: {exc}")
        return {}
    try:
        raw = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        raw = cached

    out = {}
    for sheet in cached.worksheets:
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        if sheet.title in raw.sheetnames:
            other = [list(r) for r in raw[sheet.title].iter_rows(values_only=True)]
            for i, row in enumerate(rows):
                if i >= len(other):
                    break
                for j, value in enumerate(row):
                    if value is None and j < len(other[i]):
                        rows[i][j] = other[i][j]
                        text = _norm(other[i][j])
                        if text.startswith("=") and (path, sheet.title) not in _FLAGGED:
                            _FLAGGED.add((path, sheet.title))
                            coord = f"{get_column_letter(j + 1)}{i + 1}"
                            print(f"artifacts: {sheet.title}!{coord} holds the formula "
                                  f"{text!r} and no cached value; section 8 of "
                                  "metrics/Definitions.md requires recorded values, "
                                  "not formulas")
        out[sheet.title] = rows
    return out


def _answer_sheet(workspace: Path):
    """(header row, data rows) of the sheet section 8 names, or (None, [])."""
    path = workspace / "no_contact.xlsx"
    if not path.exists():
        print("artifacts: no_contact.xlsx was not produced")
        return None, []
    grids = _grids(path)
    if SHEET not in grids:
        print(f"artifacts: no sheet named {SHEET!r}; the workbook carries "
              f"{sorted(grids)}, and section 8 names the sheet exactly")
        return None, []
    rows = grids[SHEET]
    if not rows:
        print(f"artifacts: {SHEET} is empty")
        return None, []
    header = [_norm(c) for c in rows[0]]
    if not any(c in [h.lower() for h in header] for c in SHEET_COLUMNS):
        print(f"artifacts: the first row of {SHEET} is {header!r}, which carries none "
              "of the specified column names; section 8 makes the first row the header")
        return None, []
    body = [r for r in rows[1:] if r and any(c is not None for c in r)]
    return header, body


@criterion(description="no_contact.xlsx carries the NoContact sheet, the four columns in the specified order, and data")
def sheet_and_columns(workspace: Path) -> float:
    header, body = _answer_sheet(workspace)
    if header is None:
        return 0.0
    if not body:
        print(f"artifacts: {SHEET} holds a header row and no data rows; section 8 "
              "asks for one row per customer without contact")
        return 0.0
    hits = 0
    for i, column in enumerate(SHEET_COLUMNS):
        if i < len(header) and header[i] == column:
            hits += 1
        else:
            found = header[i] if i < len(header) else "(nothing)"
            print(f"artifacts: column {i + 1} of {SHEET} is {found!r}, expected "
                  f"{column!r} — section 8 fixes the column order and names "
                  "metrics/Format_Example.xlsx the authority on spelling and capitalisation")
    return hits / len(SHEET_COLUMNS)


@criterion(description="the answer rows are sorted ascending by customer_id")
def row_order(workspace: Path) -> float:
    header, body = _answer_sheet(workspace)
    if header is None:
        return 0.0
    lowered = [h.lower() for h in header]
    if "customer_id" not in lowered:
        print(f"artifacts: {SHEET} has no customer_id column to order by")
        return 0.0
    col = lowered.index("customer_id")
    ids = [_norm(r[col]) if col < len(r) else "" for r in body]
    ids = [c for c in ids if c]
    if not ids:
        print(f"artifacts: {SHEET} holds no customer_id values; section 8 asks for "
              "one row per customer without contact, sorted ascending")
        return 0.0
    want = sorted(ids)
    hits = sum(1 for a, b in zip(ids, want) if a == b)
    if hits != len(ids):
        print(f"artifacts: {hits} of {len(ids)} rows sit where an ascending "
              "customer_id sort would put them; section 8 fixes the row order")
    return hits / len(ids)


@criterion(description="workings.json parses and carries the ten specified keys, filled in")
def workings_shape(workspace: Path) -> float:
    path = workspace / "workings.json"
    if not path.exists():
        print("workings_shape: workings.json was not produced")
        return 0.0
    try:
        blob = json.loads(path.read_text())
    except Exception as exc:
        print(f"workings_shape: workings.json is not valid JSON: {type(exc).__name__}: {exc}")
        return 0.0
    if not isinstance(blob, dict):
        print(f"workings_shape: workings.json holds a {type(blob).__name__}, not an object")
        return 0.0

    extra = [k for k in blob if k not in WORKINGS_KEYS]
    if extra:
        print(f"workings_shape: workings.json carries keys section 9 does not list: {extra}")

    hits = 0
    for key in WORKINGS_KEYS:
        if key not in blob:
            print(f"workings_shape: {key} is absent from workings.json")
            continue
        value = blob[key]
        if value is None:
            print(f"workings_shape: {key} is null; section 9 asks for the figure you derived")
            continue
        if key in WORKINGS_STRINGS:
            if isinstance(value, str) and value.strip():
                hits += 1
            else:
                print(f"workings_shape: {key} is {value!r}; section 9 asks for the literal string")
            continue
        number = value
        if isinstance(number, bool):
            number = None
        elif isinstance(number, str):
            try:
                number = float(number.strip())
            except ValueError:
                number = None
        elif not isinstance(number, (int, float)):
            number = None
        if number is None:
            print(f"workings_shape: {key} is {value!r}, which is not a number")
        elif number == 0:
            print(f"workings_shape: {key} still holds the section 9 placeholder 0 "
                  "rather than a derived figure")
        else:
            hits += 1
    return hits / max(len(WORKINGS_KEYS), len(blob))
