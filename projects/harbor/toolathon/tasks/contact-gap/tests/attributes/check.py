"""Per-customer attributes: the name, the segment and the data_issue reason code.

Two criteria. The first is the fraction of expected rows whose three graded cells
all match. The second isolates the cells the specification says are undefined and
checks the sentinel is spelled as specified — 0, an empty cell, null and a dropped
row are all rejected, because all four are common substitutes for admitting that
a value could not be established.
"""

from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/no_contact.xlsx")
COLUMNS = ("customer_id", "customer_name", "segment", "data_issue")
SENTINEL = "UNDETERMINED"

# Sheets already named in a formula diagnostic: one line per affected sheet, not
# one per cell and not one per criterion.
_FLAGGED: set = set()


def _norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _sheet(path: Path):
    import openpyxl
    from openpyxl.utils import get_column_letter

    if not path.exists():
        return None, []
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"attributes: could not open {path}: {type(exc).__name__}: {exc}")
        return None, []
    try:
        raw = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        raw = cached

    for title in sorted(cached.sheetnames, key=lambda t: (t != "NoContact", t)):
        rows = [list(r) for r in cached[title].iter_rows(values_only=True)]
        if title in raw.sheetnames:
            other = [list(r) for r in raw[title].iter_rows(values_only=True)]
            for i, row in enumerate(rows):
                if i >= len(other):
                    break
                for j, value in enumerate(row):
                    if value is None and j < len(other[i]):
                        rows[i][j] = other[i][j]
                        text = _norm(other[i][j])
                        if text.startswith("=") and (path, title) not in _FLAGGED:
                            _FLAGGED.add((path, title))
                            coord = f"{get_column_letter(j + 1)}{i + 1}"
                            print(f"attributes: {title}!{coord} holds the formula "
                                  f"{text!r} and no cached value; section 8 of "
                                  "metrics/Definitions.md requires recorded values, "
                                  "not formulas")
        for start, row in enumerate(rows[:5]):
            header = [_norm(c).lower() for c in row]
            if all(c in header for c in COLUMNS):
                idx = {c: header.index(c) for c in COLUMNS}
                body = [r for r in rows[start + 1:] if r and any(c is not None for c in r)]
                return idx, body
    print(f"attributes: no sheet in {path.name} carried all of {COLUMNS}")
    return None, []


def _records(path: Path) -> tuple[dict[str, tuple[str, str, str]], int]:
    idx, body = _sheet(path)
    if idx is None:
        return {}, 0
    out: dict[str, tuple[str, str, str]] = {}
    for row in body:
        def cell(name: str) -> str:
            i = idx[name]
            return _norm(row[i]) if i < len(row) else ""

        cid = cell("customer_id")
        if cid:
            out.setdefault(cid, (cell("customer_name"), cell("segment"), cell("data_issue")))
    return out, len(body)


@criterion(description="customer_name, segment and data_issue recorded as specified")
def attributes(workspace: Path) -> float:
    expected, _ = _records(EXPECTED)
    if not expected:
        print("attributes: the answer key could not be read")
        return 0.0
    got, rows = _records(workspace / "no_contact.xlsx")
    hits = sum(1 for cid, want in expected.items() if got.get(cid) == want)
    return hits / max(len(expected), rows, len(got))


@criterion(description="undefined attributes carry the sentinel and its reason code")
def undefined_sentinels(workspace: Path) -> float:
    expected, _ = _records(EXPECTED)
    undefined = {cid: want for cid, want in expected.items()
                 if SENTINEL in (want[0], want[1])}
    if not undefined:
        print("undefined_sentinels: the answer key holds no undefined attribute")
        return 0.0
    got, _ = _records(workspace / "no_contact.xlsx")

    hits = 0
    for cid, want in undefined.items():
        have = got.get(cid)
        if have is None:
            print(f"undefined_sentinels: {cid} is missing from the sheet; an "
                  "attribute that cannot be established is still a reported row")
            continue
        for column, wanted, actual in zip(("customer_name", "segment"), want, have):
            if wanted == SENTINEL and actual != SENTINEL:
                print(f"undefined_sentinels: {cid}.{column} is {actual!r}; the "
                      f"specification asks for {SENTINEL!r}, and 0, an empty cell "
                      "and null are all rejected")
        if have[:2] == want[:2] and have[2] == want[2]:
            hits += 1
        elif have[2] != want[2]:
            print(f"undefined_sentinels: {cid}.data_issue is {have[2]!r}, expected {want[2]!r}")
    return hits / len(undefined)
