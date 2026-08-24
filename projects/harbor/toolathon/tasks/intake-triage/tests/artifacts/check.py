"""The triage sheet exists, carries data, and is laid out the way C.8 states.

Three fractional criteria, none of which an empty shell can bank:

  ``required_columns`` - the FormID/Status/MissingItems header, looked for on every
  sheet rather than on ``.active``, standing above at least one data row. A
  workbook carrying the headers and nothing else has produced no triage, so it
  scores zero here rather than banking a third of the dimension.

  ``sheet_and_header`` - the four layout facts C.8 states: the sheet is named
  Triage, the header is row 1, the three columns are in the stated order, and
  nothing else sits on that row. Fraction of the four, not a boolean.

  ``row_order`` - C.8 asks for one row per in-scope form in ascending FormID
  order, so the order is compared positionally against the expected sheet: the
  fraction of rows standing where the answer key puts them.

Cells are read twice, ``data_only=True`` first. A workbook written by a program
that never opened Excel has no formula cache, so a formula cell reads blank
there; the formula-bearing load is the fallback, and a cell that turns out to
hold a formula is named once per sheet, with its coordinate, rather than being
reported as an empty cell.
"""

from __future__ import annotations

from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/triage.xlsx")
REQUIRED = ("formid", "status", "missingitems")
SHEET_NAME = "Triage"


def _norm(value) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "").replace("_", "")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _column_letter(position: int) -> str:
    letter = ""
    position += 1
    while position:
        position, remainder = divmod(position - 1, 26)
        letter = chr(ord("A") + remainder) + letter
    return letter


class Sheet:
    """One located sheet: its header row and the data rows under it."""

    def __init__(self, label: str, name: str, top: int, header: list, rows: list,
                 raws: list) -> None:
        self.label = label
        self.name = name
        self.top = top
        self.header = header
        self.rows = rows
        self.raws = raws
        self._reported = False

    def index(self, column: str) -> int:
        head = [_norm(c) for c in self.header]
        return head.index(column) if column in head else -1

    def _report(self, row_number: int, position: int, formula: str) -> None:
        if self._reported:
            return
        self._reported = True
        print(f"artifacts: {self.label} {self.name}!{_column_letter(position)}"
              f"{row_number} holds the formula {formula!r} and no cached value; "
              "C.9 requires recorded values, not formulas")

    def data(self) -> list[tuple[int, tuple, tuple]]:
        out: list[tuple[int, tuple, tuple]] = []
        for offset, values in enumerate(self.rows[self.top + 1:]):
            position = self.top + 1 + offset
            raw = self.raws[position] if position < len(self.raws) else ()
            values = values or ()
            raw = raw or ()
            if all(_cell(c) == "" for c in values) and all(_cell(c) == "" for c in raw):
                continue
            out.append((position + 1, values, raw))
        return out

    def value(self, row_number: int, values: tuple, raw: tuple, position: int) -> str:
        if position < 0:
            return ""
        text = _cell(values[position]) if position < len(values) else ""
        if text:
            return text
        fallback = _cell(raw[position]) if position < len(raw) else ""
        if fallback.startswith("="):
            self._report(row_number, position, fallback)
        return fallback


def _locate(path: Path, label: str) -> tuple[float, list, Sheet | None]:
    """Best header match on any sheet, and the sheet carrying a complete one."""
    import openpyxl

    if not path.exists():
        print(f"artifacts: {label} {path} was never produced")
        return 0.0, [], None
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"artifacts: could not open {label} {path}: {type(exc).__name__}: {exc}")
        return 0.0, [], None
    try:
        literal = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        literal = None

    best = 0.0
    seen: list = []
    found: Sheet | None = None
    for name in cached.sheetnames:
        rows = list(cached[name].iter_rows(values_only=True))
        raws: list = []
        if literal is not None and name in literal.sheetnames:
            raws = list(literal[name].iter_rows(values_only=True))
        for top, row in enumerate(rows[:10]):
            if row is None:
                continue
            merged = [
                cell if _cell(cell) else (raws[top][i] if top < len(raws)
                                          and raws[top] is not None
                                          and i < len(raws[top]) else None)
                for i, cell in enumerate(row)
            ]
            cells = {_norm(c) for c in merged if _cell(c)}
            if not cells:
                continue
            hit = sum(1 for c in REQUIRED if c in cells) / len(REQUIRED)
            if hit > best:
                best = hit
                seen = sorted(cells)
            if hit == 1.0 and found is None:
                found = Sheet(label, name, top, merged, rows, raws)
        if found is not None:
            break
    return best, seen, found


@criterion(description="triage.xlsx carries the FormID, Status and MissingItems "
                       "columns above at least one triaged form")
def required_columns(workspace: Path) -> float:
    best, seen, sheet = _locate(workspace / "triage.xlsx", "produced")
    if best < 1.0:
        if seen:
            print(f"required_columns: best header row found was {seen}; "
                  f"wanted {list(REQUIRED)}")
        return best
    if not sheet.data():
        print("required_columns: sheet "
              f"{sheet.name!r} carries the header and no data rows; C.8 asks for "
              "one row per in-scope form, so an empty sheet is not a partial answer")
        return 0.0
    return 1.0


@criterion(description="the sheet name and header row C.8 specifies")
def sheet_and_header(workspace: Path) -> float:
    _, _, sheet = _locate(workspace / "triage.xlsx", "produced")
    if sheet is None or not sheet.data():
        return 0.0
    header = [_cell(c) for c in sheet.header]
    while header and header[-1] == "":
        header.pop()
    facts = {
        "sheet named Triage": sheet.name.strip() == SHEET_NAME,
        "header on row 1": sheet.top == 0,
        "columns in the stated order": [_norm(c) for c in header[:3]] == list(REQUIRED),
        "nothing else on the header row": len(header) == 3,
    }
    for name, ok in facts.items():
        if not ok:
            print(f"sheet_and_header: {name} - no; sheet {sheet.name!r}, "
                  f"header row {sheet.top + 1} reads {header}")
    return sum(1 for ok in facts.values() if ok) / len(facts)


def _ids(path: Path, label: str) -> list[str]:
    _, _, sheet = _locate(path, label)
    if sheet is None:
        return []
    position = sheet.index("formid")
    out: list[str] = []
    for row_number, values, raw in sheet.data():
        out.append(sheet.value(row_number, values, raw, position).upper())
    return out


@criterion(description="rows in the ascending FormID order C.8 asks for")
def row_order(workspace: Path) -> float:
    expected = _ids(EXPECTED, "expected")
    if not expected:
        return 0.0
    got = _ids(workspace / "triage.xlsx", "produced")
    if not got:
        return 0.0
    hits = sum(1 for i, want in enumerate(expected) if i < len(got) and got[i] == want)
    if hits < len(expected):
        first = next((i for i, want in enumerate(expected)
                      if i >= len(got) or got[i] != want), 0)
        print(f"row_order: {hits} of {len(expected)} rows stand in the expected "
              f"position; first divergence at row {first + 2}, expected "
              f"{expected[first]!r}, got {got[first] if first < len(got) else '(no row)'!r}")
    return hits / max(len(expected), len(got))
