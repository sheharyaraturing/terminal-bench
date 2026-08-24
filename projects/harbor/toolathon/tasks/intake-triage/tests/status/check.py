"""Fraction of in-scope forms given the status the checklist produces.

Partial work scores: 54 of 60 statuses right reads 0.9, not 0. Rows the agent
invented — a withdrawn form triaged, a superseded revision given a row of its own —
enlarge the denominator, so padding the sheet lowers the score instead of raising it.
"""

from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/triage.xlsx")
HEADERS = ("formid", "status", "missingitems")
ABSENT = object()


def _norm_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "").replace("_", "")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _column_letter(position: int) -> str:
    letter = ""
    position += 1
    while position:
        position, remainder = divmod(position - 1, 26)
        letter = chr(ord("A") + remainder) + letter
    return letter


def _read(path: Path, label: str) -> dict[str, str]:
    """FormID -> Status, from whichever sheet carries the header.

    Loaded twice. `data_only=True` returns *cached* values, and a workbook written
    by a program that never opened Excel has no cache, so a formula cell reads
    blank there; the second, formula-bearing load is what a blank cell falls back
    to, and a cell that turns out to hold a formula is reported as such rather than
    being scored as an empty string.
    """
    import openpyxl

    if not path.exists():
        print(f"status: {label} {path} does not exist")
        return {}
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"status: could not open {label} {path}: {type(exc).__name__}: {exc}")
        return {}
    try:
        literal = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        literal = None

    for name in cached.sheetnames:
        rows = list(cached[name].iter_rows(values_only=True))
        raws: list = []
        if literal is not None and name in literal.sheetnames:
            raws = list(literal[name].iter_rows(values_only=True))
        reported = [False]
        for top, row in enumerate(rows[:10]):
            if row is None:
                continue
            head = [_norm_header(c) for c in row]
            if not all(h in head for h in HEADERS):
                continue
            idx = {h: head.index(h) for h in HEADERS}
            out: dict[str, str] = {}
            for offset, data in enumerate(rows[top + 1:]):
                raw = raws[top + 1 + offset] if top + 1 + offset < len(raws) else ()
                data = data or ()
                raw = raw or ()
                # A row written entirely as formulas is blank in the cached read and
                # is still a row: it is skipped only when the literal read is empty too.
                if all(c is None for c in data) and all(c is None for c in raw):
                    continue

                def pick(column: str) -> str:
                    position = idx[column]
                    value = _cell(data[position]) if position < len(data) else ""
                    if value:
                        return value
                    fallback = _cell(raw[position]) if position < len(raw) else ""
                    if fallback.startswith("=") and not reported[0]:
                        reported[0] = True
                        coord = f"{_column_letter(position)}{top + 2 + offset}"
                        print(f"status: {label} {name}!{coord} holds the formula "
                              f"{fallback!r} and no cached value; C.9 of the appendix "
                              "requires recorded values, not formulas")
                    return fallback

                form_id = pick("formid").upper()
                if form_id:
                    out[form_id] = pick("status").lower()
            return out

    print(f"status: no sheet in {label} {path} has a "
          f"{list(HEADERS)} header; sheets were {cached.sheetnames}")
    return {}


@criterion(description="in-scope forms carrying the status the checklist produces")
def status(workspace: Path) -> float:
    expected = _read(EXPECTED, "expected")
    if not expected:
        return 0.0
    got = _read(workspace / "triage.xlsx", "produced")
    if not got:
        return 0.0
    hits = 0
    for form_id, want in expected.items():
        if got.get(form_id, ABSENT) == want:
            hits += 1
    return hits / max(len(expected), len(got))
