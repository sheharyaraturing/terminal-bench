"""The two deliverables exist and carry the shape the definitions specify.

Shape only: whether the figures are right is graded by the segments, suppression and
workings dimensions. Scored as fractions of the required columns, the layout section 9
states and the workings keys, so a workbook that is half-shaped does not read the same
as no workbook at all — and a workbook or workings file carrying nothing but the empty
shape reads the same as no deliverable at all, which is what it is.
"""

import json
from pathlib import Path

from rewardkit import criterion

EXPECTED_DIR = Path("/tests/expected")
WORKBOOK = "conversion_by_segment.xlsx"
WORKINGS = "workings.json"

WORKINGS_KEYS = (
    "metric", "period_start", "period_end", "rows_read", "eligible_visits",
    "excluded", "segments_published", "segments_suppressed",
    "total_numerator", "total_denominator", "total_rate",
)

COLUMNS = ("Segment", "Numerator", "Denominator", "Rate", "ReasonCode")


def _grid(sheet):
    out = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                out[(cell.row, cell.column)] = cell.value
    return out


def _label(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _data_rows(path) -> int:
    """Non-empty rows below the header on the sheet that looks most like the output.

    A workbook carrying nothing but the required headers is a stub, not an answer, and
    every criterion in this dimension refuses to pay for one.
    """
    import openpyxl

    try:
        book_v = openpyxl.load_workbook(path, data_only=True)
        book_f = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        return 0
    wanted = {c.lower() for c in COLUMNS}
    best = 0
    for name in book_v.sheetnames:
        cells = dict(_grid(book_f[name])) if name in book_f.sheetnames else {}
        cells.update(_grid(book_v[name]))
        by_row: dict[int, list] = {}
        for (r, _c), value in cells.items():
            by_row.setdefault(r, []).append(value)
        header_row = next(
            (r for r in sorted(by_row)
             if {_label(v).lower() for v in by_row[r]} & wanted),
            None,
        )
        if header_row is None:
            continue
        best = max(best, sum(1 for r in by_row
                             if r > header_row and any(_label(v) for v in by_row[r])))
    return best


def read_layout(path):
    """Sheet title, header order and row order of the sheet that carries the header.

    The tables the other dimensions build are keyed by segment and so cannot see
    position; section 9 states a sheet name, a header row, a column order and a row
    order, and this is what grades all four.
    """
    import openpyxl

    if not path.exists():
        return None
    try:
        book_v = openpyxl.load_workbook(path, data_only=True)
        book_f = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"read_layout: cannot open {path.name}: {type(exc).__name__}: {exc}")
        return None

    wanted = [c.lower() for c in COLUMNS]
    for name in book_v.sheetnames:
        vals = _grid(book_v[name])
        raw = _grid(book_f[name]) if name in book_f.sheetnames else {}

        def look(r, c):
            return _label(vals.get((r, c), raw.get((r, c))))

        rows = sorted({r for r, _ in vals} | {r for r, _ in raw})
        for header_row in rows:
            labels = {look(header_row, c).lower(): c for c in range(1, 40)}
            if not all(w in labels for w in wanted):
                continue
            base = labels["segment"]
            header = [look(header_row, base + i) for i in range(len(COLUMNS))]
            hits = sum(1 for i, column in enumerate(COLUMNS)
                       if header[i].lower() == column.lower())
            body = [look(r, base).upper() for r in rows if r > header_row]
            return {
                "sheet": name,
                "header_row": header_row,
                "header": header,
                "column_hits": hits,
                "rows": [label for label in body if label],
            }
    return None


def _placeholder(blob) -> bool:
    """True when every figure in the workings file is still section 10's zero."""

    def numbers(node):
        if isinstance(node, bool):
            return
        if isinstance(node, (int, float)):
            yield float(node)
        elif isinstance(node, dict):
            for value in node.values():
                yield from numbers(value)
        elif isinstance(node, list):
            for value in node:
                yield from numbers(value)

    found = list(numbers(blob))
    return bool(found) and all(value == 0.0 for value in found)


@criterion(description="conversion_by_segment.xlsx carries the five required columns over data")
def required_columns(workspace: Path) -> float:
    import openpyxl

    path = workspace / WORKBOOK
    if not path.exists():
        return 0.0
    try:
        book = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"required_columns: cannot open {WORKBOOK}: {type(exc).__name__}: {exc}")
        return 0.0
    best = 0
    for name in book.sheetnames:
        seen = {_label(c.value).lower() for row in book[name].iter_rows() for c in row}
        best = max(best, sum(1 for c in COLUMNS if c.lower() in seen))
    if best == 0:
        return 0.0
    if _data_rows(path) == 0:
        print(f"required_columns: {WORKBOOK} carries headers and no data row; "
              "section 9 requires a row per segment in segments.csv and a TOTAL row")
        return 0.0
    return best / len(COLUMNS)


@criterion(description="the sheet name, header row, column order and row order of section 9")
def output_layout(workspace: Path) -> float:
    key = read_layout(EXPECTED_DIR / WORKBOOK)
    if key is None:
        print("output_layout: the answer key is unreadable")
        return 0.0
    got = read_layout(workspace / WORKBOOK)
    if got is None or not got["rows"]:
        print(f"output_layout: no sheet of {WORKBOOK} carries the five columns over at "
              "least one data row")
        return 0.0

    sheet_score = 1.0 if got["sheet"].strip().lower() == key["sheet"].strip().lower() else 0.0
    if sheet_score == 0.0:
        print(f"output_layout: the table sits on sheet {got['sheet']!r}; "
              f"section 9 names {key['sheet']!r}")

    header_score = 1.0 if got["header_row"] == key["header_row"] else 0.0
    if header_score == 0.0:
        print(f"output_layout: the header is on row {got['header_row']}; "
              "section 9 puts it on row 1")

    column_score = got["column_hits"] / len(COLUMNS)
    if column_score < 1.0:
        print(f"output_layout: columns read {got['header']}; "
              f"section 9 states {list(COLUMNS)} in that order")

    want, mine = key["rows"], got["rows"]
    span = max(len(want), len(mine))
    placed = sum(1 for i, label in enumerate(want) if i < len(mine) and mine[i] == label)
    row_score = placed / span
    if placed < span:
        print(f"output_layout: {placed} of {span} rows sit where section 9 puts them "
              f"(ascending segment_id, then TOTAL); read {mine}")

    return (sheet_score + header_score + column_score + row_score) / 4


@criterion(description="workings.json parses and carries the keys the definitions name")
def workings_shape(workspace: Path) -> float:
    path = workspace / WORKINGS
    if not path.exists():
        return 0.0
    try:
        blob = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        print(f"workings_shape: {WORKINGS} is not readable JSON: {type(exc).__name__}: {exc}")
        return 0.0
    if not isinstance(blob, dict):
        print(f"workings_shape: {WORKINGS} holds {type(blob).__name__}, not an object")
        return 0.0
    if _placeholder(blob):
        print(f"workings_shape: every figure in {WORKINGS} is still 0 — this is section 10's "
              "placeholder, not a workings file")
        return 0.0
    return sum(1 for k in WORKINGS_KEYS if k in blob) / len(WORKINGS_KEYS)
