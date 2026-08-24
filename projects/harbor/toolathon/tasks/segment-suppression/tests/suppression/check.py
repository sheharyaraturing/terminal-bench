"""The suppression policy, graded on its own.

Three independent things, because they fail independently:

1. an undefined rate is the literal sentinel — 0, 0.00, an empty cell, null, a missing
   row and the thin sample's own rate are each rejected explicitly;
2. the two undefined causes carry their own reason codes and are not interchanged;
3. the total row aggregates the published segments only.
"""

import json
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from rewardkit import criterion

EXPECTED_DIR = Path("/tests/expected")
DIMENSION = "suppression"
WORKBOOK = "conversion_by_segment.xlsx"
WORKINGS = "workings.json"

COLUMNS = ("Segment", "Numerator", "Denominator", "Rate", "ReasonCode")
SENTINEL = "UNDEFINED"
REASONS = ("OK", "LOW_BASE", "NO_BASE")

_REF = re.compile(r"\$?([A-Z]{1,3})\$?([0-9]{1,7})")
_RANGE = re.compile(
    r"SUM\(\s*\$?([A-Z]{1,3})\$?([0-9]{1,7})\s*:\s*\$?([A-Z]{1,3})\$?([0-9]{1,7})\s*\)",
    re.IGNORECASE,
)
_SAFE = re.compile(r"^[0-9eE+\-*/(). ,]*$")


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _round_half_up(value, places=0):
    q = Decimal(1).scaleb(-int(places))
    return float(Decimal(repr(float(value))).quantize(q, rounding=ROUND_HALF_UP))


def _resolve(raw, vals, row, col, depth):
    """Value of one cell, evaluating a formula when the cached value is absent.

    A formula outside the evaluated subset comes back as its own text at the top
    level, so the failure reads as the formula that could not be evaluated rather
    than as an empty cell. Nested lookups keep returning None, which is what the
    arithmetic above expects.
    """
    got = vals.get((row, col))
    if got is not None:
        return got
    formula = raw.get((row, col))
    if isinstance(formula, str) and formula.startswith("="):
        value = _evaluate(formula[1:], raw, vals, depth + 1)
        return formula if value is None and depth == 0 else value
    return formula


def _evaluate(expr, raw, vals, depth=0):
    """Evaluate the small formula subset a spreadsheet tool is likely to emit."""
    if depth > 8:
        return None

    def as_number(value):
        if isinstance(value, bool) or value is None:
            return 0.0
        if isinstance(value, (int, float, Decimal)):
            return float(value)
        try:
            return float(str(value).strip().rstrip("%"))
        except ValueError:
            raise ValueError("non-numeric cell in formula")

    try:
        def sum_range(m):
            c1, r1, c2, r2 = _col_index(m.group(1)), int(m.group(2)), _col_index(m.group(3)), int(m.group(4))
            total = 0.0
            for r in range(min(r1, r2), max(r1, r2) + 1):
                for c in range(min(c1, c2), max(c1, c2) + 1):
                    total += as_number(_resolve(raw, vals, r, c, depth))
            return repr(total)

        text = _RANGE.sub(sum_range, expr)
        text = _REF.sub(
            lambda m: repr(as_number(_resolve(raw, vals, int(m.group(2)), _col_index(m.group(1)), depth))),
            text,
        )
    except ValueError:
        return None

    guarded = re.sub(r"ROUND|ABS", "", text, flags=re.IGNORECASE)
    if not _SAFE.match(guarded):
        return None
    try:
        return eval(  # noqa: S307 - characters are whitelisted above
            text, {"__builtins__": {}}, {"ROUND": _round_half_up, "round": _round_half_up, "ABS": abs}
        )
    except Exception:
        return None


_FORMULA_SEEN = set()


def _letters(col: int) -> str:
    out = ""
    while col:
        col, rest = divmod(col - 1, 26)
        out = chr(65 + rest) + out
    return out


def _report_formula(path, sheet, coord, value) -> None:
    """Name a formula cell that reached grading as its own text.

    Once per sheet and not once per cell, so a workbook written entirely in formulas
    does not bury the run's real failures under one line per cell.
    """
    key = (str(path), sheet)
    if key in _FORMULA_SEEN:
        return
    _FORMULA_SEEN.add(key)
    print(f"{DIMENSION}: {sheet}!{coord} holds the formula {value!r} and no cached value, "
          f"and it is not one this verifier can evaluate; Definitions.md section 9 requires "
          f"recorded values, not formulas")


def _grid(sheet):
    out = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                out[(cell.row, cell.column)] = cell.value
    return out


def _label(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_table(path):
    """Segment -> row dict, from whichever sheet carries the header. {} when unreadable.

    Two loads: the cached-value read, then the formula-bearing read as a fallback, so a
    workbook whose cells are formulas written outside Excel is still gradeable.
    """
    import openpyxl

    if not path.exists():
        return {}
    try:
        book_v = openpyxl.load_workbook(path, data_only=True)
        book_f = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"read_table: cannot open {path.name}: {type(exc).__name__}: {exc}")
        return {}

    wanted = [c.lower() for c in COLUMNS]
    for name in book_v.sheetnames:
        vals = _grid(book_v[name])
        raw = _grid(book_f[name]) if name in book_f.sheetnames else {}
        rows = sorted({r for r, _ in vals} | {r for r, _ in raw})
        header_row = None
        index = {}
        for r in rows:
            labels = {_label(vals.get((r, c), raw.get((r, c)))).lower(): c
                      for c in range(1, 40)}
            if all(w in labels for w in wanted):
                header_row = r
                index = {w: labels[w] for w in wanted}
                break
        if header_row is None:
            continue

        table = {}
        stranded = None
        for r in rows:
            if r <= header_row:
                continue
            cells = {w: _resolve(raw, vals, r, c, 0) for w, c in index.items()}
            if stranded is None:
                for w, value in cells.items():
                    if isinstance(value, str) and value.startswith("="):
                        stranded = (f"{_letters(index[w])}{r}", value)
                        break
            key = _label(cells["segment"]).upper()
            if not key:
                continue
            table[key] = cells
        if stranded is not None:
            _report_formula(path, name, *stranded)
        return table

    print(f"read_table: no sheet in {path.name} carries the header {COLUMNS}")
    return {}


def as_int(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        return int(value) if float(value).is_integer() else None
    text = str(value).strip().replace(",", "")
    try:
        return int(text) if float(text).is_integer() else None
    except ValueError:
        return None


def as_rate(value):
    """The Rate cell as ('number', float) or ('text', str) or ('missing', '')."""
    if value is None:
        return ("missing", "")
    if isinstance(value, bool):
        return ("text", str(value))
    if isinstance(value, (int, float, Decimal)):
        return ("number", float(value))
    text = str(value).strip()
    if not text:
        return ("missing", "")
    try:
        return ("number", float(text.rstrip("%").strip()))
    except ValueError:
        return ("text", text.upper())


TOTAL = "TOTAL"


def _rejected(value) -> str:
    """Name the near-miss form of undefined this cell is, if it is one."""
    kind, payload = as_rate(value)
    if kind == "missing":
        return "an empty cell" if value is not None else "no value"
    if kind == "number":
        return "the number 0" if payload == 0.0 else f"the number {payload}"
    if payload in ("NULL", "NONE", "NAN", "N/A", "NA", "#DIV/0!"):
        return f"the placeholder {payload!r}"
    return ""


@criterion(description="undefined rates carry the sentinel, not 0, blank or null")
def undefined_rate_sentinel(workspace: Path) -> float:
    expected = read_table(EXPECTED_DIR / WORKBOOK)
    if not expected:
        print("undefined_rate_sentinel: the answer key is unreadable")
        return 0.0
    undefined = {k: v for k, v in expected.items()
                 if as_rate(v["rate"]) == ("text", SENTINEL)}
    if not undefined:
        print("undefined_rate_sentinel: the answer key holds no undefined rate")
        return 0.0

    got = read_table(workspace / WORKBOOK)
    hits = 0
    for key in undefined:
        row = got.get(key)
        if row is None:
            print(f"undefined_rate_sentinel: {key} has no row; a missing row is not the sentinel")
            continue
        kind, payload = as_rate(row["rate"])
        if kind == "text" and payload == SENTINEL:
            hits += 1
        else:
            why = _rejected(row["rate"]) or f"the text {payload!r}"
            print(f"undefined_rate_sentinel: {key} rate is {why}, want {SENTINEL!r}")
    return hits / len(undefined)


@criterion(description="LOW_BASE and NO_BASE are not interchanged")
def reason_codes(workspace: Path) -> float:
    expected = read_table(EXPECTED_DIR / WORKBOOK)
    if not expected:
        return 0.0
    got = read_table(workspace / WORKBOOK)
    if not got:
        return 0.0
    hits = 0
    for key, want in expected.items():
        row = got.get(key)
        if row is None:
            continue
        wanted = _label(want["reasoncode"]).upper()
        mine = _label(row["reasoncode"]).upper()
        if mine == wanted:
            hits += 1
        elif {mine, wanted} == {"LOW_BASE", "NO_BASE"}:
            print(f"reason_codes: {key} reports {mine}; the two undefined causes are swapped")
        else:
            print(f"reason_codes: {key} want {wanted!r} got {mine!r}")
    return hits / max(len(expected), len(got))


@criterion(description="the total row aggregates the published segments only")
def total_excludes_suppressed(workspace: Path) -> float:
    expected = read_table(EXPECTED_DIR / WORKBOOK)
    if not expected or TOTAL not in expected:
        print("total_excludes_suppressed: the answer key has no TOTAL row")
        return 0.0
    got = read_table(workspace / WORKBOOK)
    row = got.get(TOTAL)
    if row is None:
        print("total_excludes_suppressed: the produced workbook has no TOTAL row")
        return 0.0
    want = expected[TOTAL]

    hits = 0
    for field in ("numerator", "denominator"):
        if as_int(row[field]) is not None and as_int(row[field]) == as_int(want[field]):
            hits += 1
        else:
            print(f"total_excludes_suppressed: TOTAL.{field} want {want[field]!r} "
                  f"got {row[field]!r} — suppressed segments contribute nothing")
    kind_w, value_w = as_rate(want["rate"])
    kind_g, value_g = as_rate(row["rate"])
    if kind_w == kind_g and (
        (kind_w == "number" and abs(value_g - value_w) <= 1e-6) or
        (kind_w == "text" and value_g == value_w)
    ):
        hits += 1
    else:
        print(f"total_excludes_suppressed: TOTAL.rate want {want['rate']!r} got {row['rate']!r}")
    return hits / 3
