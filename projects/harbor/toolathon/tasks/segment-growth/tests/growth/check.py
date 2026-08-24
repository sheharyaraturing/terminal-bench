"""All computable growth cells must match the final v2 answer key."""

from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/growth.xlsx")
NA = "N/A"


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        text = value.strip()
        if text.upper().replace(" ", "") in ("N/A", "NA", "#N/A"):
            return NA
        if "%" in text:
            return text
        try:
            return f"{float(text.replace(',', '')):.1f}"
        except ValueError:
            return text
    try:
        x = float(value)
        if x == 0:
            x = 0.0
        return f"{x:.1f}"
    except (TypeError, ValueError):
        return str(value).strip()


def _read(path: Path):
    import openpyxl

    if not path.exists():
        return None
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"could not open {path.name}: {type(exc).__name__}: {exc}")
        return None
    if wb.sheetnames != ["Growth"]:
        return None
    sh = wb["Growth"]
    if str(sh.cell(1, 1).value or "").strip() != "Segment":
        return None

    years = []
    for col in range(2, sh.max_column + 1):
        value = sh.cell(1, col).value
        if value is None:
            return None
        try:
            years.append(str(int(float(str(value).strip()))))
        except ValueError:
            years.append(str(value).strip())

    order, table = [], {}
    for row in range(2, sh.max_row + 1):
        value = sh.cell(row, 1).value
        if value is None:
            continue
        segment = str(value).strip()
        if segment in table:
            return None
        order.append(segment)
        table[segment] = {
            year: _cell(sh.cell(row, i + 2).value)
            for i, year in enumerate(years)
        }
    return years, order, table


@criterion(description="every computable growth figure matches the methodology and exact rounding")
def growth(workspace: Path) -> float:
    want = _read(EXPECTED)
    got = _read(workspace / "growth.xlsx")
    if want is None or got is None:
        return 0.0
    want_years, want_order, want_table = want
    got_years, _, got_table = got
    if got_years != want_years:
        return 0.0

    wrong = []
    for segment in want_order:
        for year in want_years:
            expected = want_table[segment][year]
            if expected == NA:
                continue
            actual = got_table.get(segment, {}).get(year, "")
            if actual != expected:
                wrong.append(f"{segment} {year}: got {actual or '(blank)'!r}, expected {expected!r}")
    if wrong:
        print(f"{len(wrong)} computable growth cell(s) incorrect: {'; '.join(wrong[:6])}")
        return 0.0
    return 1.0
