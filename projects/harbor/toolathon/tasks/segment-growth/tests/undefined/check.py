"""Every expected undefined growth cell must be reported as N/A."""

from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/growth.xlsx")
NA = "N/A"


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        text = value.strip()
        if text.upper().replace(" ", "") in ("N/A", "NA", "#N/A"):
            return NA
        return text
    return str(value).strip()


def _read(path: Path):
    import openpyxl

    if not path.exists():
        return None
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"could not open {path.name}: {type(exc).__name__}: {exc}")
        return None
    if wb.sheetnames != ["Growth"]:
        return None
    sh = wb["Growth"]
    if str(sh.cell(1, 1).value or "").strip() != "Segment":
        return None

    years = []
    for col in range(2, sh.max_column + 1):
        value = sh.cell(1, col).value
        if value is None:
            return None
        try:
            years.append(str(int(float(str(value).strip()))))
        except ValueError:
            years.append(str(value).strip())

    order, table = [], {}
    for row in range(2, sh.max_row + 1):
        value = sh.cell(row, 1).value
        if value is None:
            continue
        segment = str(value).strip()
        if segment in table:
            return None
        order.append(segment)
        table[segment] = {
            year: _cell(sh.cell(row, i + 2).value)
            for i, year in enumerate(years)
        }
    return years, order, table


@criterion(description="every undefined growth cell is explicitly reported as N/A")
def undefined(workspace: Path) -> float:
    want = _read(EXPECTED)
    got = _read(workspace / "growth.xlsx")
    if want is None or got is None:
        return 0.0
    want_years, want_order, want_table = want
    got_years, _, got_table = got
    if got_years != want_years:
        return 0.0

    wrong = []
    for segment in want_order:
        for year in want_years:
            if want_table[segment][year] != NA:
                continue
            actual = got_table.get(segment, {}).get(year, "")
            if actual != NA:
                wrong.append(f"{segment} {year}: got {actual or '(blank)'!r}, expected 'N/A'")
    if wrong:
        print(f"{len(wrong)} undefined cell(s) incorrect: {'; '.join(wrong[:6])}")
        return 0.0
    return 1.0
