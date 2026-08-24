"""Exact required year columns and canonical segment row order."""

from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/growth.xlsx")


def _shape(path: Path):
    import openpyxl

    if not path.exists():
        return None
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"could not open {path.name}: {type(exc).__name__}: {exc}")
        return None
    if wb.sheetnames != ["Growth"]:
        return None
    sh = wb["Growth"]
    if str(sh.cell(1, 1).value or "").strip() != "Segment":
        return None

    years = []
    for col in range(2, sh.max_column + 1):
        value = sh.cell(1, col).value
        if value is None:
            print(f"blank year header in column {col}")
            return None
        try:
            years.append(str(int(float(str(value).strip()))))
        except ValueError:
            years.append(str(value).strip())

    order = []
    seen = set()
    for row in range(2, sh.max_row + 1):
        value = sh.cell(row, 1).value
        if value is None:
            # Match the legacy verifier: rows with no Segment label do not define a segment.
            continue
        segment = str(value).strip()
        if segment in seen:
            print(f"duplicate segment row: {segment}")
            return None
        seen.add(segment)
        order.append(segment)
    return years, order


@criterion(description="year columns and segment rows match the expected canonical shape exactly")
def shape(workspace: Path) -> float:
    want = _shape(EXPECTED)
    got = _shape(workspace / "growth.xlsx")
    if want is None or got is None:
        return 0.0
    if got[0] != want[0]:
        print(f"year columns are {got[0]}, expected {want[0]}")
        return 0.0
    if got[1] != want[1]:
        print(f"segment row order differs; got {got[1][:5]}, expected {want[1][:5]}")
        return 0.0
    return 1.0
