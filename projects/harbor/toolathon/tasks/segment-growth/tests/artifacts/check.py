"""Strict workbook/representation contract for growth.xlsx."""

from pathlib import Path

from rewardkit import criterion


@criterion(description="growth.xlsx exists and obeys the literal-value workbook contract")
def artifact_contract(workspace: Path) -> float:
    import openpyxl

    path = workspace / "growth.xlsx"
    if not path.exists():
        print(f"missing expected output: {path}")
        return 0.0

    try:
        wb_values = openpyxl.load_workbook(path, data_only=True)
        wb_formulas = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"could not open growth.xlsx: {type(exc).__name__}: {exc}")
        return 0.0

    if wb_values.sheetnames != ["Growth"] or wb_formulas.sheetnames != ["Growth"]:
        print(f"growth.xlsx must contain exactly one sheet named 'Growth'; got {wb_values.sheetnames}")
        return 0.0

    sheet = wb_values["Growth"]
    formula_sheet = wb_formulas["Growth"]
    if sheet.cell(1, 1).value is None or str(sheet.cell(1, 1).value).strip() != "Segment":
        print("growth.xlsx first header must be 'Segment'")
        return 0.0

    for row in formula_sheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                print(f"Excel formula found at {cell.coordinate}; literal computed values are required")
                return 0.0
            if isinstance(cell.value, str) and "%" in cell.value:
                print(f"percent string found at {cell.coordinate}; use a numeric percentage-point value")
                return 0.0
            if cell.value is not None and isinstance(cell.number_format, str) and "%" in cell.number_format:
                print(f"percent number format found at {cell.coordinate}; use the direct numeric growth value")
                return 0.0

    return 1.0
