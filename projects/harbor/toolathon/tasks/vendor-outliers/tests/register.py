"""Reading and scoring for exception_register.xlsx, shared by every dimension.

Root-level criteria are imported before the subdirectories and are marked
``shared=True`` so they register against the dimension that calls them rather
than against a reward of their own. Each subdirectory's ``check.py`` is then
one line: the dimension is the choice of criteria, not a copy of the reader.
"""

from __future__ import annotations

import datetime
import re
from pathlib import Path
from typing import NamedTuple

from rewardkit import criterion

REGISTER = "exception_register.xlsx"
EXPECTED = Path(__file__).resolve().parent / "expected" / REGISTER

# The layout Format_Example.xlsx fixes, in the order it fixes it.
REQUIRED = ("vendor", "month", "amount", "baseline", "deviation", "status")
# Money columns; the policy reports these to two decimal places.
ROUNDED = ("baseline", "deviation")
TOTAL_LABEL = "total flagged:"
MONTH_RE = re.compile(r"^(\d{4})-(\d{1,2})(?:-\d{1,2})?$")


def _norm(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _num(value) -> str:
    """Two-decimal text so 1200 and 1200.00 compare equal; N/A stays literal."""
    try:
        return f"{float(str(value).replace(',', '')):.2f}"
    except (TypeError, ValueError):
        return _norm(value).upper()


def _month(value) -> str:
    """Normalise a Month cell to YYYY-MM.

    Excel hands back a datetime for anything it parsed as a date, and an agent
    may write 2025-6 as readily as 2025-06. They all name the same month.
    """
    if isinstance(value, (datetime.datetime, datetime.date)):
        return f"{value.year:04d}-{value.month:02d}"
    text = _norm(value)
    match = MONTH_RE.match(text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"
    return text


def _numeric(value) -> bool:
    try:
        float(str(value).replace(",", ""))
        return True
    except (TypeError, ValueError):
        return False


def _is_rounded(value) -> bool:
    """True when value carries no more than two decimal places.

    Non-numeric cells (N/A) are not a rounding question, so they pass.
    """
    try:
        number = float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return True
    return abs(number - round(number, 2)) < 1e-9


def _trim(row) -> list:
    """Drop trailing empty cells so a padded row compares like a tight one."""
    cells = list(row)
    while cells and _norm(cells[-1]) == "":
        cells.pop()
    return cells


class Register(NamedTuple):
    header: list[str]  # the header row as written, trimmed and lowercased
    entries: dict  # (vendor, month) -> (amount, baseline, deviation, status)
    total: str | None  # the cell beside 'Total flagged:', or None if absent
    total_extra: list[str]  # anything else sharing the footer row
    duplicates: list  # vendor-months listed more than once
    unrounded: list[str]  # money reported to more than two decimals


def read(path: Path) -> Register | None:
    """Parse a register, or None if there is nothing readable to parse.

    The sheet is located by its header rather than by ``.active``: an agent may
    leave a scratch sheet selected, and that is not what is being graded here.
    Columns are read by name so a reordered header still yields rows -- the
    order itself is scored separately, by ``register_layout``.
    """
    import openpyxl

    if not path.exists():
        print(f"{path.name} is missing")
        return None
    try:
        book = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"could not open {path.name}: {type(exc).__name__}: {exc}")
        return None

    for title in book.sheetnames:
        rows = list(book[title].iter_rows(values_only=True))
        if not rows or rows[0] is None:
            continue
        header = [_norm(c).lower() for c in _trim(rows[0])]
        if any(column not in header for column in REQUIRED):
            continue
        idx = {column: header.index(column) for column in REQUIRED}

        entries, total, extra = {}, None, []
        duplicates, unrounded = [], []
        for row in rows[1:]:
            if row is None or all(c is None for c in row):
                continue
            if _norm(row[0]).lower() == TOTAL_LABEL:
                total = _norm(row[1]) if len(row) > 1 else ""
                # Keep whatever else sits on this row: a footer written over the
                # last data row instead of below it is worth naming precisely.
                extra = [_norm(c) for c in row[2:] if _norm(c)]
                continue
            vendor = _norm(row[idx["vendor"]])
            if not vendor:
                continue
            month = _month(row[idx["month"]])
            key = (vendor.lower(), month)
            if key in entries:
                duplicates.append(key)
            for column in ROUNDED:
                if not _is_rounded(row[idx[column]]):
                    unrounded.append(f"{vendor} {month} {column}={_norm(row[idx[column]])}")
            entries[key] = (
                _num(row[idx["amount"]]),
                _num(row[idx["baseline"]]),
                _num(row[idx["deviation"]]),
                _norm(row[idx["status"]]).upper(),
            )
        return Register(header, entries, total, extra, duplicates, unrounded)

    print(f"no sheet in {path.name} carries the columns {list(REQUIRED)}; "
          f"sheets: {book.sheetnames}")
    return None


def _scored(expected: dict, got: dict) -> float:
    """Fraction of expected rows reproduced exactly. Extra rows count against the
    denominator, so padding cannot raise the score."""
    if not expected:
        return 1.0 if not got else 0.0
    hits = sum(1 for key, value in expected.items() if got.get(key) == value)
    return hits / max(len(expected), len(got))


@criterion(description="exception_register.xlsx carries the six required columns",
           shared=True)
def register_columns(workspace: Path) -> float:
    """Fraction of the six columns present on some sheet -- partial credit for a
    register that is nearly the right shape."""
    import openpyxl

    path = workspace / REGISTER
    if not path.exists():
        return 0.0
    try:
        book = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return 0.0
    best = 0.0
    for title in book.sheetnames:
        rows = list(book[title].iter_rows(values_only=True))
        if not rows or rows[0] is None:
            continue
        header = {_norm(c).lower() for c in rows[0] if c is not None}
        best = max(best, sum(1 for c in REQUIRED if c in header) / len(REQUIRED))
    return best


@criterion(description="the columns are in the order Format_Example.xlsx fixes",
           shared=True)
def register_layout(workspace: Path) -> bool:
    register = read(workspace / REGISTER)
    if register is None:
        return False
    if register.header == list(REQUIRED):
        return True
    print(f"column layout does not match Format_Example.xlsx: "
          f"expected {list(REQUIRED)}, found {register.header}")
    return False


@criterion(description="each vendor-month is listed once", shared=True)
def register_unique(workspace: Path) -> bool:
    register = read(workspace / REGISTER)
    if register is None:
        return False
    if not register.duplicates:
        return True
    listed = sorted({f"{vendor} {month}" for vendor, month in register.duplicates})
    print(f"{len(register.duplicates)} vendor-month(s) listed more than once: {listed[:5]}")
    return False


@criterion(description="baseline and deviation are reported to two decimal places",
           shared=True)
def register_rounded(workspace: Path) -> bool:
    register = read(workspace / REGISTER)
    if register is None:
        return False
    if not register.unrounded:
        return True
    print(f"{len(register.unrounded)} figure(s) carry more than two decimal places — "
          f"{'; '.join(register.unrounded[:3])}")
    return False


@criterion(description="rows reported as {status}, with baseline and deviation",
           shared=True)
def register_rows(workspace: Path, status: str) -> float:
    expected = read(EXPECTED)
    if expected is None:
        raise FileNotFoundError(f"ground truth register unreadable: {EXPECTED}")
    register = read(workspace / REGISTER)
    got = register.entries if register else {}
    want = {k: v for k, v in expected.entries.items() if v[3] == status}
    have = {k: v for k, v in got.items() if v[3] == status}
    missing = sorted(set(want) - set(have))
    unexpected = sorted(set(have) - set(want))
    if missing:
        print(f"{status}: missing {missing[:5]}")
    if unexpected:
        print(f"{status}: unexpected {unexpected[:5]}")
    for key in sorted(set(want) & set(have)):
        if have[key] != want[key]:
            print(f"{status}: {key[0]} {key[1]} reads {have[key]}, expected {want[key]}")
    return _scored(want, have)


@criterion(description="the Total flagged: row holds the number of FLAGGED rows",
           shared=True)
def register_total(workspace: Path) -> bool:
    expected = read(EXPECTED)
    if expected is None:
        raise FileNotFoundError(f"ground truth register unreadable: {EXPECTED}")
    register = read(workspace / REGISTER)
    if register is None:
        return False
    if register.total is None:
        print(f"no row starting {TOTAL_LABEL!r}")
        return False
    if _num(register.total) == _num(expected.total):
        return True
    # Same verdict either way; the wording is what changes. A non-numeric cell
    # here usually means the footer landed on top of a data row.
    if _numeric(register.total):
        print(f"'Total flagged:' reads {register.total!r}, expected {expected.total!r}")
    else:
        detail = (f"the cell beside {TOTAL_LABEL!r} holds {register.total!r}, which is not "
                  f"a count (expected {expected.total!r})")
        if register.total_extra:
            detail += (f"; that row also carries {register.total_extra[:4]}, so the footer "
                       f"looks written over the last data row rather than appended below it")
        print(detail)
    return False
