"""Fraction of the UnmatchedLedger sheet's graded cells that agree with the answer key.

The reason column is graded alongside the identifying ones: an entry left unsettled
for the wrong stated reason has not been reconciled, only listed.
"""
import re
from pathlib import Path

EXPECTED = Path("/tests/expected/reconciliation.xlsx")
ARTIFACT = "reconciliation.xlsx"
DIMENSION = "unmatched_ledger"


def _read(path: Path, sheet_name: str, columns: list) -> list:
    """Rows of the named sheet, or of any sheet whose header carries every column.

    The workbook is opened twice: once for cached values and once for what the cells
    actually hold. A workbook written outside Excel stores no cached result, so a formula
    cell reads blank on the first pass; the second pass surfaces the formula text, which
    is then reported by name and compared as written. Section 8.6 of the procedure
    requires recorded values, so a formula is a wrong answer, not a blank one.
    """
    import openpyxl

    if not path.exists():
        return []
    try:
        data = openpyxl.load_workbook(path, data_only=True)
        raw = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"could not open {path}: {type(exc).__name__}: {exc}")
        return []

    wanted = {c.strip().lower() for c in columns}

    def header_of(sheet):
        rows = sheet.iter_rows(values_only=True)
        try:
            first = next(rows)
        except StopIteration:
            return []
        return [str(c).strip().lower() if c is not None else "" for c in first]

    chosen = None
    for name in data.sheetnames:
        if name.strip().lower() == sheet_name.strip().lower():
            chosen = name
            break
    if chosen is None or not wanted.issubset(set(header_of(data[chosen]))):
        for name in data.sheetnames:
            if wanted.issubset(set(header_of(data[name]))):
                chosen = name
                break
    if chosen is None:
        print(f"{path.name}: no sheet carries the {sheet_name} columns {sorted(wanted)}")
        return []

    view, formulas = data[chosen], raw[chosen]
    header = header_of(view)
    missing = [c for c in wanted if c not in header]
    if missing:
        print(f"{path.name}/{chosen}: header is missing {sorted(missing)}")
        return []
    index = {c.strip().lower(): header.index(c.strip().lower()) for c in columns}

    out = []
    reported = False
    for number, row in enumerate(view.iter_rows(values_only=True), start=1):
        if number == 1 or row is None or all(c is None for c in row):
            continue
        record = {}
        for column in columns:
            position = index[column.strip().lower()]
            value = row[position] if position < len(row) else None
            if value is None:
                cell = formulas.cell(row=number, column=position + 1)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    if not reported:
                        print(f"{DIMENSION}: {chosen}!{cell.coordinate} holds the formula "
                              f"{cell.value!r} and no cached value; section 8.6 of the "
                              f"procedure requires recorded values, not formulas")
                        reported = True
                value = cell.value
            record[column] = value
        out.append(record)
    return out


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip().upper()


def _date(value) -> str:
    from datetime import date, datetime

    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    return f"{match.group(1)}-{match.group(2)}-{match.group(3)}" if match else text


def _number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", "").replace("GBP", "").replace("\u00a3", "").strip()
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return float(text)
    except ValueError:
        return None


def _ids(value) -> str:
    """Sections 5.5 and 8.2 fix both the separator and the ascending order; keep both."""
    parts = [p.strip().upper() for p in str(value if value is not None else "").split(";")]
    return "|".join(p for p in parts if p)


def _same(kind, want, got) -> bool:
    if kind == "number":
        a, b = _number(want), _number(got)
        return a is not None and b is not None and abs(a - b) < 0.005
    if kind == "date":
        return _date(want) == _date(got)
    if kind == "ids":
        return _ids(want) == _ids(got)
    return _text(want) == _text(got)


def _score(sheet_name, key, columns, workspace):
    want = _read(EXPECTED, sheet_name, [key] + [c for c, _ in columns])
    if not want:
        print(f"{sheet_name}: answer key unreadable; scoring 0")
        return 0.0
    got = _read(Path(workspace) / ARTIFACT, sheet_name, [key] + [c for c, _ in columns])
    got_by_key = {}
    for row in got:
        got_by_key.setdefault(_text(row[key]), row)
    hits = 0
    for row in want:
        other = got_by_key.get(_text(row[key]))
        if other is None:
            continue
        hits += sum(1 for column, kind in columns if _same(kind, row[column], other[column]))
    denominator = max(len(want), len(got_by_key)) * len(columns)
    return hits / denominator if denominator else 0.0


from rewardkit import criterion

COLUMNS = [
    ("Vendor", "text"),
    ("Invoice_Ref", "text"),
    ("Due_Date", "date"),
    ("Amount_Due", "number"),
    ("Reason", "text"),
]


@criterion(description="ledger entries left unsettled, every column compared")
def unmatched_ledger(workspace: Path) -> float:
    return _score("UnmatchedLedger", "Ledger_ID", COLUMNS, workspace)
