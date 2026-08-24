"""The reconciliation workbook exists, carries the shape the procedure fixes, and holds data.

Three criteria, all fractional and all conditioned on the sheet actually carrying rows:

* `required_columns` -- the three sheets are present, each with its column names and at
  least one data row. A workbook that carries nothing but the headers scores zero here,
  so producing an empty shell earns nothing.
* `sheet_and_column_order` -- section 8.1 fixes the sheet order and the column order; both
  are compared positionally, one point per sheet position and per column position.
* `row_order` -- section 8.1 fixes the row sort of each sheet; the produced key column is
  compared position by position against the answer key's.
"""

import re
from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/reconciliation.xlsx")
ARTIFACT = "reconciliation.xlsx"
DIMENSION = "artifacts"

SHEETS = ["Matched", "UnmatchedBank", "UnmatchedLedger"]
SHAPE = {
    "Matched": ["Txn_ID", "Value_Date", "Ledger_IDs", "Vendor", "Bank_Amount",
                "Ledger_Amount", "Amount_Difference", "Date_Gap_Days", "Match_Type"],
    "UnmatchedBank": ["Txn_ID", "Value_Date", "Counterparty", "Reference", "Amount",
                      "Reason", "Candidate_Ledger_IDs"],
    "UnmatchedLedger": ["Ledger_ID", "Vendor", "Invoice_Ref", "Due_Date", "Amount_Due",
                        "Reason"],
}
KEYS = {"Matched": "Txn_ID", "UnmatchedBank": "Txn_ID", "UnmatchedLedger": "Ledger_ID"}


def _open(path: Path):
    import openpyxl

    if not path.exists():
        print(f"{DIMENSION}: {path.name} was never produced")
        return None, None
    try:
        return (openpyxl.load_workbook(path, data_only=True),
                openpyxl.load_workbook(path, data_only=False))
    except Exception as exc:
        print(f"{DIMENSION}: could not open {path}: {type(exc).__name__}: {exc}")
        return None, None


def _header(sheet) -> list:
    try:
        first = next(sheet.iter_rows(values_only=True))
    except StopIteration:
        return []
    return [str(c).strip().lower() if c is not None else "" for c in first]


def _locate(book, sheet_name: str, columns: list):
    """The sheet of that name, or any sheet whose header carries every column."""
    wanted = {c.strip().lower() for c in columns}
    for name in book.sheetnames:
        if name.strip().lower() == sheet_name.strip().lower():
            if wanted.issubset(set(_header(book[name]))):
                return name
            break
    for name in book.sheetnames:
        if wanted.issubset(set(_header(book[name]))):
            return name
    return None


def _rows(view, formulas, position: int, sheet_label: str) -> list:
    """Values of one column, skipping the header. Formula cells are reported, not guessed."""
    out = []
    reported = False
    for number, row in enumerate(view.iter_rows(values_only=True), start=1):
        if number == 1 or row is None or all(c is None for c in row):
            continue
        value = row[position] if position < len(row) else None
        if value is None and formulas is not None:
            cell = formulas.cell(row=number, column=position + 1)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                if not reported:
                    print(f"{DIMENSION}: {sheet_label}!{cell.coordinate} holds the formula "
                          f"{cell.value!r} and no cached value; section 8.6 of the procedure "
                          f"requires recorded values, not formulas")
                    reported = True
                value = cell.value
            else:
                value = cell.value
        out.append(value)
    return out


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip().upper()


def _data_rows(sheet) -> int:
    count = 0
    for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if number == 1 or row is None or all(c is None for c in row):
            continue
        count += 1
    return count


@criterion(description="the three sheets carry their columns and at least one data row")
def required_columns(workspace: Path) -> float:
    book, _ = _open(Path(workspace) / ARTIFACT)
    if book is None:
        return 0.0

    hits = 0
    total = 0
    for sheet, columns in SHAPE.items():
        total += len(columns)
        found = _locate(book, sheet, columns)
        if found is None:
            print(f"{DIMENSION}: no sheet named {sheet} and none carrying its columns")
            continue
        if _data_rows(book[found]) == 0:
            print(f"{DIMENSION}: sheet {found} carries the {sheet} headers and no data rows; "
                  f"a header-only sheet is not a reconciliation")
            continue
        header = set(_header(book[found]))
        hits += sum(1 for c in columns if c.lower() in header)
    return hits / total


@criterion(description="sheet order and column order follow section 8.1")
def sheet_and_column_order(workspace: Path) -> float:
    book, _ = _open(Path(workspace) / ARTIFACT)
    if book is None:
        return 0.0

    names = [n.strip().lower() for n in book.sheetnames]
    hits = 0
    total = len(SHEETS)
    for index, sheet in enumerate(SHEETS):
        # A sheet counts only when it is in the right place and carries rows: an empty
        # sheet in the right position is not evidence of a reconciliation.
        if (index < len(names) and names[index] == sheet.lower()
                and _data_rows(book[book.sheetnames[index]]) > 0):
            hits += 1
        else:
            actual = book.sheetnames[index] if index < len(names) else "(absent)"
            print(f"{DIMENSION}: sheet {index + 1} is {actual!r} where section 8.1 puts "
                  f"{sheet} carrying rows")

    for sheet, columns in SHAPE.items():
        total += len(columns)
        found = _locate(book, sheet, columns)
        if found is None or _data_rows(book[found]) == 0:
            continue
        header = _header(book[found])
        for position, column in enumerate(columns):
            if position < len(header) and header[position] == column.lower():
                hits += 1
    return hits / total


@criterion(description="each sheet is sorted as section 8.1 requires")
def row_order(workspace: Path) -> float:
    want_book, _ = _open(EXPECTED)
    if want_book is None:
        print(f"{DIMENSION}: answer key unreadable; scoring 0")
        return 0.0
    got_book, got_formulas = _open(Path(workspace) / ARTIFACT)
    if got_book is None:
        return 0.0

    hits = 0
    total = 0
    for sheet in SHEETS:
        columns = SHAPE[sheet]
        key = KEYS[sheet]
        want_sheet = _locate(want_book, sheet, columns)
        want = _rows(want_book[want_sheet], None,
                     _header(want_book[want_sheet]).index(key.lower()), sheet)
        got_sheet = _locate(got_book, sheet, columns)
        if got_sheet is None:
            total += len(want)
            continue
        header = _header(got_book[got_sheet])
        got = _rows(got_book[got_sheet], got_formulas[got_sheet],
                    header.index(key.lower()), got_sheet)
        total += max(len(want), len(got))
        hits += sum(1 for a, b in zip(want, got) if _text(a) == _text(b))
    return hits / total if total else 0.0
