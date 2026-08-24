"""Per-customer revenue, graded field by field against the answer key.

Scores fractions. A run that gets 10 of 11 customers right reads 0.91, not 0. Rows
the agent invented count against the denominator, so padding the sheet cannot lift
the score, and a customer the definition says to leave out costs a row when it is
present.

Every column the definition fixes is compared, `Undefined_Reason` included: §7 says it
carries `mixed_currency` where revenue is undefined and is left blank otherwise, so a
run that fills the blank cells with anything at all is wrong there.

The undefined customer is scored separately, because `0`, an empty cell, `null` and
a missing row are the four ways a model usually gets that one wrong and each of them
has to fail.

§7 also fixes the layout — the sheet is named `Revenue`, the six columns run in a stated
order, and the rows run by `Customer_ID` ascending. Stating an order and then keying rows
by id leaves the requirement unmeasured, so both orders are compared positionally here,
each as a fraction rather than a pass/fail.
"""

from decimal import Decimal, InvalidOperation
from pathlib import Path

import rewardkit as rk
from rewardkit import criterion

EXPECTED = Path("/tests/expected/revenue_by_customer.xlsx")
COLUMNS = (
    "customer_id",
    "customer_name",
    "currency",
    "orders",
    "revenue",
    "undefined_reason",
)
FIELDS = ("customer_name", "currency", "orders", "revenue", "undefined_reason")
SENTINEL = "undefined"
CENT = Decimal("0.01")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().lower()


def _money(value):
    """Decimal at two places, or None when the cell is not a number."""
    text = _text(value).replace(",", "").replace("$", "").replace("€", "")
    if not text:
        return None
    try:
        return Decimal(text).quantize(CENT)
    except (InvalidOperation, ValueError):
        return None


class Sheet:
    """One parsed workbook: the rows in file order, plus the layout §7 fixes."""

    def __init__(self, name: str, header: list[str], rows: list[dict]) -> None:
        self.name = name
        self.header = header
        self.rows = rows
        self.ids = [_text(r["customer_id"]).upper() for r in rows]
        self.by_id = {}
        for row in rows:
            key = _text(row["customer_id"]).upper()
            if key:
                self.by_id[key] = row

    def __bool__(self) -> bool:
        return bool(self.by_id)


_PARSED: dict[str, Sheet] = {}


def _read(path: Path, label: str) -> Sheet:
    """The sheet that carries the header, keyed by id and kept in file order.

    Parsed once per file and cached, so the reader's own diagnostics — a missing file,
    a missing header, a sheet with no rows, a cell still holding a formula — are
    printed once for the run rather than once per criterion.

    Loaded twice on purpose. `data_only=True` returns *cached* values, and a
    workbook written by a program that never opened Excel has no cache, so every
    formula cell reads back as None. The second load supplies what the first
    dropped; a cell still holding a formula is named as a formula — once per sheet,
    so a wholly formula-written answer does not flood the log — rather than being
    reported as a blank the agent never wrote.
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    cached_parse = _PARSED.get(str(path))
    if cached_parse is not None:
        return cached_parse

    empty = Sheet("", [], [])
    if not path.exists():
        print(f"{label}: {path.name} was never written")
        return _PARSED.setdefault(str(path), empty)
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
        raw = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"{label}: could not open {path.name}: {type(exc).__name__}: {exc}")
        return _PARSED.setdefault(str(path), empty)

    for name in cached.sheetnames:
        cached_rows = [list(r) for r in cached[name].iter_rows(values_only=True)]
        raw_rows = [list(r) for r in raw[name].iter_rows(values_only=True)]
        for top, row in enumerate(cached_rows[:5]):
            header = [_text(c) for c in row]
            if not all(col in header for col in COLUMNS):
                continue
            index = {col: header.index(col) for col in COLUMNS}
            out: list[dict] = []
            reported = False
            for offset, cached_row in enumerate(cached_rows[top + 1:], start=top + 1):
                raw_row = raw_rows[offset] if offset < len(raw_rows) else []
                merged = [
                    cell if cell is not None else (raw_row[i] if i < len(raw_row) else None)
                    for i, cell in enumerate(cached_row)
                ]
                if all(c is None for c in merged):
                    continue
                record = {}
                for col, i in index.items():
                    value = merged[i] if i < len(merged) else None
                    if isinstance(value, str) and value.startswith("="):
                        if not reported:
                            coord = f"{get_column_letter(i + 1)}{offset + 1}"
                            print(
                                f"{label}: {name}!{coord} holds the formula {value!r} and no "
                                f"cached value; §7 of metrics/Definitions.md requires recorded "
                                f"values, not formulas — further formula cells on this sheet "
                                f"are not listed"
                            )
                            reported = True
                        value = None
                    record[col] = value
                out.append(record)
            sheet = Sheet(name, header, out)
            if sheet:
                return _PARSED.setdefault(str(path), sheet)
            print(f"{label}: sheet {name!r} has the header but no data rows")
            return _PARSED.setdefault(str(path), empty)
    print(f"{label}: no sheet in {path.name} carries the six required columns")
    return _PARSED.setdefault(str(path), empty)


def _matches(field: str, want, got) -> bool:
    if field == "revenue":
        if _text(want) == SENTINEL:
            return _text(got) == SENTINEL
        want_money, got_money = _money(want), _money(got)
        return want_money is not None and want_money == got_money
    if field == "orders":
        try:
            return int(Decimal(_text(want))) == int(Decimal(_text(got)))
        except (InvalidOperation, ValueError):
            return False
    return _text(want) == _text(got)


@criterion(
    description="customer name, currency, order count and revenue match the definition",
    shared=True,
)
def revenue_fields(workspace: Path) -> float:
    expected = _read(EXPECTED, "revenue_fields[key]")
    if not expected:
        return 0.0
    got = _read(workspace / "revenue_by_customer.xlsx", "revenue_fields")

    hits = 0
    for cid, want in expected.by_id.items():
        have = got.by_id.get(cid)
        if have is None:
            print(f"revenue_fields: no row for {cid}")
            continue
        for field in FIELDS:
            if _matches(field, want[field], have.get(field)):
                hits += 1
            else:
                print(
                    f"revenue_fields: {cid} {field} expected {want[field]!r}, "
                    f"got {have.get(field)!r}"
                )
    return hits / (len(FIELDS) * max(len(expected.by_id), len(got.by_id)))


rk.revenue_fields(weight=3.0)


@criterion(
    description="the mixed-currency customer is reported undefined, not resolved to a number",
)
def undefined_customer(workspace: Path) -> float:
    expected = _read(EXPECTED, "undefined_customer[key]")
    if not expected:
        return 0.0
    flagged = {
        cid: want
        for cid, want in expected.by_id.items()
        if _text(want["revenue"]) == SENTINEL
    }
    if not flagged:
        print("undefined_customer: the answer key has no undefined row; the task lost its restraint case")
        return 0.0
    got = _read(workspace / "revenue_by_customer.xlsx", "undefined_customer")

    hits = 0
    for cid, want in flagged.items():
        have = got.by_id.get(cid)
        if have is None:
            print(f"undefined_customer: {cid} is missing; an undefined row is still a row")
            continue
        if _text(have.get("revenue")) == SENTINEL:
            hits += 1
        else:
            print(
                f"undefined_customer: {cid} revenue is {have.get('revenue')!r}; the "
                f"definition wants the sentinel, and 0, blank and null are all wrong"
            )
        if _text(have.get("undefined_reason")) == _text(want["undefined_reason"]):
            hits += 1
        else:
            print(
                f"undefined_customer: {cid} reason expected "
                f"{want['undefined_reason']!r}, got {have.get('undefined_reason')!r}"
            )
    return hits / (2 * len(flagged))


@criterion(
    description="sheet name and column order, as §7 fixes them",
    shared=True,
)
def sheet_layout(workspace: Path) -> float:
    expected = _read(EXPECTED, "sheet_layout[key]")
    if not expected:
        return 0.0
    got = _read(workspace / "revenue_by_customer.xlsx", "sheet_layout")
    if not got:
        return 0.0

    want_header = [c for c in expected.header if c]
    have_header = [c for c in got.header if c]
    placed = sum(
        1
        for i, col in enumerate(want_header)
        if i < len(have_header) and have_header[i] == col
    )
    columns = placed / max(len(want_header), len(have_header))
    if columns < 1.0:
        print(
            f"sheet_layout: columns expected in the order {want_header}, "
            f"got {have_header} — {placed} of {len(want_header)} in place"
        )

    named = 1.0 if got.name.strip().lower() == expected.name.strip().lower() else 0.0
    if not named:
        print(f"sheet_layout: §7 names the sheet {expected.name!r}; the data sits on {got.name!r}")
    return (named + columns) / 2


rk.sheet_layout(weight=0.5)


@criterion(
    description="rows run by Customer_ID ascending, as §7 fixes them",
    shared=True,
)
def row_order(workspace: Path) -> float:
    expected = _read(EXPECTED, "row_order[key]")
    if not expected:
        return 0.0
    got = _read(workspace / "revenue_by_customer.xlsx", "row_order")
    if not got:
        return 0.0

    placed = sum(
        1 for i, cid in enumerate(expected.ids) if i < len(got.ids) and got.ids[i] == cid
    )
    total = max(len(expected.ids), len(got.ids))
    if placed < total:
        print(
            f"row_order: §7 orders rows by Customer_ID ascending; {placed} of {total} "
            f"rows sit in the right position (expected {expected.ids}, got {got.ids})"
        )
    return placed / total


rk.row_order(weight=0.5)
