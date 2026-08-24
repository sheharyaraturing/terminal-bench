"""Both artifacts appeared, carry the shape the definition document fixes, and hold data.

This dimension only asks whether the deliverables are well formed. Whether the
figures are right is `tests/revenue/` and `tests/workings/`.

Shape alone is never enough. A workbook holding the six headers and no rows, and a
`workings.json` holding the §8 skeleton with its zeros still in it, are the two ways an
agent can produce a file-shaped nothing; neither is an answer and neither scores here.
Existence is folded into these two criteria rather than graded on its own, so a run that
touches two empty files banks nothing.
"""

import json
from pathlib import Path

from rewardkit import criterion

SHEET_COLUMNS = (
    "customer_id",
    "customer_name",
    "currency",
    "orders",
    "revenue",
    "undefined_reason",
)

WORKINGS_KEYS = (
    "metric",
    "value",
    "undefined",
    "undefined_reason",
    "denominator",
    "population",
)

POPULATION_KEYS = ("included", "excluded", "rows_before_join", "rows_after_join")

EXCLUDED_KEYS = ("internal_account", "out_of_period", "cancelled_or_draft")


def _text(value) -> str:
    return "" if value is None else str(value).strip().lower()


@criterion(
    description="revenue_by_customer.xlsx carries the six columns the definition fixes, above at least one data row"
)
def sheet_columns(workspace: Path) -> float:
    import openpyxl
    from openpyxl.utils import get_column_letter

    path = workspace / "revenue_by_customer.xlsx"
    if not path.exists():
        print("sheet_columns: revenue_by_customer.xlsx was never written")
        return 0.0
    try:
        # Two reads: cached values first, then the formula-bearing one. A workbook
        # written by a program that never opened Excel has no value cache, so a
        # formula cell reads back as None from the first load alone.
        cached = openpyxl.load_workbook(path, data_only=True)
        raw = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"sheet_columns: could not open {path.name}: {type(exc).__name__}: {exc}")
        return 0.0

    # Never trust .active: an agent may leave an empty sheet in front of the real one.
    best = 0.0
    for name in cached.sheetnames:
        cached_rows = [list(r) for r in cached[name].iter_rows(values_only=True)]
        raw_rows = [list(r) for r in raw[name].iter_rows(values_only=True)]
        reported = False
        for top, row in enumerate(cached_rows[:5]):
            header = {_text(c) for c in row if c is not None}
            hit = sum(1 for c in SHEET_COLUMNS if c in header) / len(SHEET_COLUMNS)
            if hit == 0.0:
                continue

            # A formula cell is content, not a blank, so the row census reads the
            # formula-bearing load and names the first offender once per sheet.
            body = 0
            for offset in range(top + 1, max(len(cached_rows), len(raw_rows))):
                cached_row = cached_rows[offset] if offset < len(cached_rows) else []
                raw_row = raw_rows[offset] if offset < len(raw_rows) else []
                merged = [
                    cell if cell is not None else (raw_row[i] if i < len(raw_row) else None)
                    for i, cell in enumerate(cached_row)
                ] or list(raw_row)
                if all(c is None or _text(c) == "" for c in merged):
                    continue
                body += 1
                if not reported:
                    for i, cell in enumerate(merged):
                        if isinstance(cell, str) and cell.startswith("="):
                            coord = f"{get_column_letter(i + 1)}{offset + 1}"
                            print(
                                f"sheet_columns: {name}!{coord} holds the formula {cell!r} "
                                f"and no cached value; §7 of metrics/Definitions.md requires "
                                f"recorded values, not formulas"
                            )
                            reported = True
                            break
            if body == 0:
                print(
                    f"sheet_columns: sheet {name!r} carries the header but no data row; "
                    f"§7 asks for one row per customer with an in-scope order"
                )
                continue
            best = max(best, hit)
    if best < 1.0:
        print(f"sheet_columns: no sheet carries the full header over data; best match {best:.2f}")
    return best


@criterion(description="workings.json parses, carries the required keys and reports a population")
def workings_shape(workspace: Path) -> float:
    path = workspace / "workings.json"
    if not path.exists():
        print("workings_shape: workings.json was never written")
        return 0.0
    try:
        doc = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        print(f"workings_shape: workings.json is not valid JSON: {type(exc).__name__}: {exc}")
        return 0.0
    if not isinstance(doc, dict):
        print(f"workings_shape: workings.json holds a {type(doc).__name__}, not an object")
        return 0.0

    wanted = list(WORKINGS_KEYS) + [f"population.{k}" for k in POPULATION_KEYS]
    wanted += [f"population.excluded.{k}" for k in EXCLUDED_KEYS]

    population = doc.get("population")
    population = population if isinstance(population, dict) else {}
    excluded = population.get("excluded")
    excluded = excluded if isinstance(excluded, dict) else {}

    # The §8 code block is a template with zeros in it. Copying it back unfilled
    # reports no orders at all, which §3 forbids — `included` plus the three excluded
    # counts must equal the 38 rows of data/orders.csv — so it is a stub, not a record.
    counts = [population.get(k) for k in ("included", "rows_before_join", "rows_after_join")]
    counts += [excluded.get(k) for k in EXCLUDED_KEYS]
    if all(c in (0, 0.0, None, "0", "") for c in counts):
        print(
            "workings_shape: workings.json reports no orders anywhere in 'population'; "
            "that is the §8 template, not a workings record"
        )
        return 0.0

    hits = 0
    for key in WORKINGS_KEYS:
        hits += 1 if key in doc else 0
    for key in POPULATION_KEYS:
        hits += 1 if key in population else 0
    for key in EXCLUDED_KEYS:
        hits += 1 if key in excluded else 0

    if hits < len(wanted):
        missing = [k for k in WORKINGS_KEYS if k not in doc]
        missing += [f"population.{k}" for k in POPULATION_KEYS if k not in population]
        missing += [f"population.excluded.{k}" for k in EXCLUDED_KEYS if k not in excluded]
        print(f"workings_shape: missing key(s) {missing}")
    return hits / len(wanted)
