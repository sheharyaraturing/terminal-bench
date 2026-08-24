"""Per-account grading of the fee run.

Four fractions, each hits / max(len(expected), len(got)) so padding the sheet
cannot inflate a score and an empty sheet reads zero:

  balance_cents      the recorded balance restated in whole minor units
  fee_cents          the progressive fee, rounded once, half-up
  effective_rate     fee over balance to six decimal places
  undefined_sentinel the accounts whose fee is undefined carry the sentinel

The workbook is opened twice. openpyxl's data_only read returns the value Excel
cached beside a formula, and a workbook written by a program has no such cache, so
a formula cell reads as None; the second read recovers what is actually in the
cell. That does not evaluate the formula — openpyxl has no formula
engine — but it turns the failure from "the agent left the cell blank" into
a named cell holding a named formula, which is the difference between a five
minute diagnosis and a wasted afternoon. One line per affected sheet, so a
workbook written entirely in formulas reports once rather than 168 times.

Sheets are searched by header, never by .active: an agent that leaves an empty
first sheet in front of its answer is still correct. The sheet name, the column
order and the row order are graded by tests/artifacts.
"""

from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/fee_schedule_run.xlsx")
COLUMNS = ("account", "balance", "fee", "effectiverate")
SENTINEL = "UNDEFINED"
SIX = Decimal("0.000001")

_CACHE: dict[str, list[dict]] = {}


def _norm(value) -> str:
    return "" if value is None else str(value).strip().lower().replace(" ", "")


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.strip().startswith("=")


def _read(path: Path) -> list[dict]:
    """Every data row of the first sheet whose header carries the four columns."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    if not path.exists():
        print(f"fees: {path.name} is absent")
        return []
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
        literal = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"fees: could not open {path}: {type(exc).__name__}: {exc}")
        return []

    for sheet in cached.worksheets:
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        row0, col0 = sheet.min_row, sheet.min_column
        try:
            alt = [list(r) for r in literal[sheet.title].iter_rows(values_only=True)]
        except Exception:
            alt = []
        for i, row in enumerate(rows):
            header = [_norm(c) for c in row]
            if not all(c in header for c in COLUMNS):
                continue
            idx = {c: header.index(c) for c in COLUMNS}
            out = []
            formulas: list[tuple[str, str]] = []
            for j in range(i + 1, len(rows)):
                merged = list(rows[j])
                for k, value in enumerate(merged):
                    if value is None and j < len(alt) and k < len(alt[j]):
                        merged[k] = alt[j][k]
                if all(c is None for c in merged):
                    continue
                for column, k in idx.items():
                    if k < len(merged) and _is_formula(merged[k]):
                        coord = f"{get_column_letter(col0 + k)}{row0 + j}"
                        formulas.append((coord, str(merged[k]).strip()))
                cell = {c: merged[idx[c]] if idx[c] < len(merged) else None
                        for c in COLUMNS}
                account = "" if cell["account"] is None else str(cell["account"]).strip()
                if account:
                    cell["account"] = account
                    out.append(cell)
            if formulas:
                coord, text = formulas[0]
                more = (f" ({len(formulas) - 1} further formula cells on this sheet "
                        "are reported the same way)") if len(formulas) > 1 else ""
                print(f"fees: {sheet.title}!{coord} holds the formula {text!r} and no "
                      "cached value; section 10 of the fee schedule requires recorded "
                      f"values, not formulas{more}")
            return out
    print(f"fees: no sheet in {path.name} has a header row carrying {COLUMNS}")
    return []


def _table(path: Path) -> list[dict]:
    key = str(path)
    if key not in _CACHE:
        _CACHE[key] = _read(path)
    return _CACHE[key]


def _int(value):
    """A whole number of minor units, the sentinel, or None."""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if text.upper() == SENTINEL:
            return SENTINEL
        if _is_formula(text):
            return None
        text = text.replace(",", "").replace("$", "").strip()
        if not text:
            return None
        try:
            number = Decimal(text)
        except InvalidOperation:
            return None
    elif isinstance(value, bool):
        return None
    elif isinstance(value, (int, float)):
        number = Decimal(str(value))
    else:
        return None
    if number != number.to_integral_value():
        return None
    return int(number)


def _rate(value):
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if text.upper() == SENTINEL:
            return SENTINEL
        if _is_formula(text):
            return None
        text = text.replace("%", "").strip()
        if not text:
            return None
        try:
            number = Decimal(text)
        except InvalidOperation:
            return None
    elif isinstance(value, bool):
        return None
    elif isinstance(value, (int, float)):
        number = Decimal(str(value))
    else:
        return None
    return number.quantize(SIX, rounding=ROUND_HALF_UP)


def _column(path: Path, column: str, reader) -> dict:
    return {row["account"]: reader(row[column]) for row in _table(path)}


def _fraction(expected: dict, got: dict) -> float:
    if not expected:
        return 0.0
    hits = sum(1 for key, value in expected.items() if got.get(key) == value)
    return hits / max(len(expected), len(got))


@criterion(description="balances restated in whole minor units")
def balance_cents(workspace: Path) -> float:
    return _fraction(
        _column(EXPECTED, "balance", _int),
        _column(workspace / "fee_schedule_run.xlsx", "balance", _int),
    )


@criterion(description="progressive fee per account, in whole minor units")
def fee_cents(workspace: Path) -> float:
    return _fraction(
        _column(EXPECTED, "fee", _int),
        _column(workspace / "fee_schedule_run.xlsx", "fee", _int),
    )


@criterion(description="effective rate per account, to six decimal places")
def effective_rate(workspace: Path) -> float:
    return _fraction(
        _column(EXPECTED, "effectiverate", _rate),
        _column(workspace / "fee_schedule_run.xlsx", "effectiverate", _rate),
    )


@criterion(description="accounts with no defined fee carry the sentinel, not 0, blank or a dropped row")
def undefined_sentinel(workspace: Path) -> float:
    expected = {a: v for a, v in _column(EXPECTED, "fee", _int).items() if v == SENTINEL}
    if not expected:
        return 0.0
    got_all = _column(workspace / "fee_schedule_run.xlsx", "fee", _int)
    got = {a: v for a, v in got_all.items() if v == SENTINEL}
    hits = 0
    for account in expected:
        value = got_all.get(account, "<row absent>")
        if value == SENTINEL:
            hits += 1
        else:
            print(f"fees: {account} has no defined fee; the sheet says {value!r}")
    return hits / max(len(expected), len(got))
