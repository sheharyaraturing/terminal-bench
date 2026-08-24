"""Per-transaction restatement: the money, its provenance, the source fields, the order.

Four fractions rather than one boolean, so 58 of 60 rows right reads differently
from nothing at all, and each fraction names a different failure. Extra rows the
agent invented count against the denominator, so padding cannot lift the score.

Section 7.1 fixes the row order, so the row order is compared positionally rather
than by dictionary lookup -- a stated-but-ungraded ordering is a requirement nobody
measures. It stays a fraction of the rows sitting in the right place, never a boolean.

Undefined is a sentinel: only the literal string UNDEFINED counts. 0, "", null and a
missing row are all rejected, because all four are common model outputs and all four
are wrong.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import rewardkit as rk
from rewardkit import criterion

EXPECTED = Path("/tests/expected/usd_revenue.xlsx")
SHEET = "USD_Revenue"
COLUMNS = ("txn_id", "currency", "amount", "rate_date", "rate", "usd_amount", "rate_basis")
CENT = Decimal("0.01")
EIGHT = Decimal("0.00000001")


# ---------------------------------------------------------------- workbook reading
LITERAL = re.compile(r"^=\s*[+-]?\d+(?:\.\d+)?\s*$")
_ANNOUNCED: set = set()  # (workbook label, sheet) already reported: one line per sheet


def _resolve(values_sheet, formula_sheet, dimension: str) -> dict:
    """Cell grid in which an uncached formula surfaces as its formula text, not as a blank.

    load_workbook(data_only=True) returns *cached* results, and a workbook written by
    a program that never ran Excel has no cache, so every formula cell reads None.
    There is no way to evaluate the formula here -- openpyxl cannot, and a verifier
    with no network has no formula engine to add -- and section 7.1 of the policy
    states that computed values are recorded rather than formulas, so a formula-bearing
    cell breaks a stated rule and scoring it wrong is enforcement rather than a trap.

    What this does is make the failure legible: load twice, and where the cached read
    is blank, carry the formula text through, so the mismatch reads as
    "'=(C2-B2)/B2*100'" rather than as "(blank)" and nobody is sent after a phantom
    empty cell. A cell holding nothing but a literal (`=2418.63`) is recovered, because
    that needs no engine. One diagnostic line per sheet, so a fully formula-written
    answer names its cause without flooding the log.
    """
    from openpyxl.utils import get_column_letter

    grid: dict = {}
    for row in values_sheet.iter_rows():
        for cell in row:
            grid[(cell.row, cell.column)] = cell.value

    seen = (dimension, formula_sheet.title)
    announced = seen in _ANNOUNCED
    for row in formula_sheet.iter_rows():
        for cell in row:
            text = cell.value
            if not isinstance(text, str) or not text.startswith("="):
                continue
            if grid.get((cell.row, cell.column)) is not None:
                continue  # Excel cached a result and the data_only read already has it
            if LITERAL.match(text):
                grid[(cell.row, cell.column)] = Decimal(text[1:].strip())
                continue
            grid[(cell.row, cell.column)] = text
            if not announced:
                announced = True
                _ANNOUNCED.add(seen)
                coord = f"{get_column_letter(cell.column)}{cell.row}"
                print(f"{dimension}: {formula_sheet.title}!{coord} holds the formula {text!r} "
                      f"and no cached value; section 7.1 of FX_Policy.md requires recorded "
                      f"values, not formulas. Further formula cells on this sheet are not "
                      f"listed.")
    return grid


def _rows(path: Path, label: str) -> list:
    """Every data row of the deliverable, in sheet order, as a list of dicts."""
    import openpyxl

    if not path.exists():
        print(f"conversion: {label} usd_revenue.xlsx was not produced")
        return []
    try:
        values = openpyxl.load_workbook(path, data_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        print(f"conversion: could not open {label} {path}: {type(exc).__name__}: {exc}")
        return []

    # the policy names the sheet, so try that one first; fall back to a scan of the rest
    # so a misnamed sheet still has its money graded (the name is scored in `artifacts`)
    order = sorted(values.sheetnames,
                   key=lambda n: (n.strip() != SHEET, values.sheetnames.index(n)))
    for name in order:  # never trust .active
        grid = _resolve(values[name], formulas[name], f"conversion ({label})")
        width = max((c for _, c in grid), default=0)
        header = [str(grid.get((1, c)) or "").strip().lower() for c in range(1, width + 1)]
        if any(col not in header for col in COLUMNS):
            continue
        idx = {col: header.index(col) + 1 for col in COLUMNS}
        height = max((r for r, _ in grid), default=1)
        out: list = []
        for r in range(2, height + 1):
            record = {col: grid.get((r, idx[col])) for col in COLUMNS}
            if all(v is None or str(v).strip() == "" for v in record.values()):
                continue
            out.append(record)
        return out

    print(f"conversion: no sheet in the {label} workbook carries all of {COLUMNS}")
    return []


# ---------------------------------------------------------------- normalisation
def _undefined(value) -> bool:
    return isinstance(value, str) and value.strip().upper() == "UNDEFINED"


def _decimal(value, places: Decimal):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if not text:
            return None
        try:
            return Decimal(text).quantize(places, rounding=ROUND_HALF_UP)
        except InvalidOperation:
            return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(places, rounding=ROUND_HALF_UP)
        except InvalidOperation:
            return None
    return None


def _iso(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        text = value.strip()
        return text[:10] if re.match(r"^\d{4}-\d{2}-\d{2}", text) else None
    return None


def _text(value):
    return str(value).strip().upper() if value is not None else None


def _same_money(want, got) -> bool:
    if _undefined(want):
        return _undefined(got)
    return _decimal(want, CENT) is not None and _decimal(want, CENT) == _decimal(got, CENT)


def _same_rate(want, got) -> bool:
    if _undefined(want):
        return _undefined(got)
    return _decimal(want, EIGHT) is not None and _decimal(want, EIGHT) == _decimal(got, EIGHT)


def _same_date(want, got) -> bool:
    if _undefined(want):
        return _undefined(got)
    return _iso(want) is not None and _iso(want) == _iso(got)


def _keyed(rows: list[dict]) -> dict[str, dict]:
    return {_text(r["txn_id"]): r for r in rows if _text(r["txn_id"])}


def _score(workspace: Path, fields) -> float:
    want = _keyed(_rows(EXPECTED, "expected"))
    if not want:
        print("conversion: the answer key could not be read; scoring zero")
        return 0.0
    got = _keyed(_rows(workspace / "usd_revenue.xlsx", "agent"))
    hits = 0
    for txn_id, row in want.items():
        theirs = got.get(txn_id)
        if theirs is not None and all(fn(row[col], theirs[col]) for col, fn in fields):
            hits += 1
    return hits / max(len(want), len(got), 1)


@criterion(description="USD_Amount restated at the rate in force on the transaction date", shared=True)
def usd_amount(workspace: Path) -> float:
    return _score(workspace, [("usd_amount", _same_money)])


rk.usd_amount(weight=2.0)


@criterion(description="the quote actually used: its date, its value and its basis", shared=True)
def rate_provenance(workspace: Path) -> float:
    return _score(workspace, [
        ("rate_date", _same_date),
        ("rate", _same_rate),
        ("rate_basis", lambda a, b: _text(a) == _text(b)),
    ])


rk.rate_provenance(weight=2.0)


@criterion(description="source currency and source amount carried through unchanged", shared=True)
def source_fields(workspace: Path) -> float:
    return _score(workspace, [
        ("currency", lambda a, b: _text(a) == _text(b)),
        ("amount", _same_money),
    ])


rk.source_fields(weight=0.5)


@criterion(description="rows in the Txn_ID order section 7.1 fixes")
def row_order(workspace: Path) -> float:
    want = [t for t in (_text(r["txn_id"]) for r in _rows(EXPECTED, "expected")) if t]
    if not want:
        print("conversion: the answer key could not be read; scoring zero")
        return 0.0
    got = [t for t in (_text(r["txn_id"]) for r in _rows(workspace / "usd_revenue.xlsx", "agent")) if t]
    hits = sum(1 for i, txn in enumerate(want) if i < len(got) and got[i] == txn)
    if hits < max(len(want), len(got)):
        first = next(i for i, txn in enumerate(want) if i >= len(got) or got[i] != txn) \
            if hits < len(want) else len(want)
        theirs = got[first] if first < len(got) else "(no row)"
        print(f"row_order: {hits} of {len(want)} rows sit where section 7.1 puts them "
              f"({len(got)} rows written); first divergence at data row {first + 1}, "
              f"expected {want[first] if first < len(want) else '(no row)'!r}, got {theirs!r}")
    return hits / max(len(want), len(got), 1)
