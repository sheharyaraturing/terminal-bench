"""The two deliverables exist, carry the shape the policy fixes, and carry data.

Structural only: sheet name, column names, column order and JSON keys, never their
values. Nothing here pays for an untouched workspace, and nothing here pays for an
empty shell either -- every criterion is gated on the artifact carrying at least one
data row, so a workbook holding nothing but the correct header scores zero on all of
them rather than banking a third of the reward for free.
"""

from pathlib import Path

from rewardkit import criterion

SHEET = "USD_Revenue"
COLUMNS = ("txn_id", "currency", "amount", "rate_date", "rate", "usd_amount", "rate_basis")
KEYS = (
    "metric", "period_start", "period_end", "total_usd_revenue",
    "transactions_by_currency", "carried_forward_count", "undefined_count",
    "undefined_by_reason",
)


def _populated_sheets(workspace: Path, dimension: str):
    """Every sheet that carries a header row and at least one data row under it.

    Returns [(sheet_name, header_cells)]. Never trusts `.active` -- an agent that
    leaves an empty leading Sheet1 and writes to the second sheet is correct.
    """
    import openpyxl

    path = workspace / "usd_revenue.xlsx"
    if not path.exists():
        print(f"{dimension}: usd_revenue.xlsx was not produced")
        return None
    try:
        book = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"{dimension}: could not open {path}: {type(exc).__name__}: {exc}")
        return None

    out = []
    for name in book.sheetnames:
        rows = list(book[name].iter_rows(values_only=True))
        if not rows or rows[0] is None:
            continue
        header = [c for c in rows[0]]
        if not any(c is not None and str(c).strip() for c in header):
            continue
        body = [r for r in rows[1:]
                if r is not None and any(c is not None and str(c).strip() != "" for c in r)]
        if not body:
            continue
        out.append((name, header))
    if not out:
        print(f"{dimension}: no sheet in usd_revenue.xlsx carries a header and at least one "
              f"data row; section 7.1 asks for one row per transaction in the population")
        return None
    return out


@criterion(description="usd_revenue.xlsx carries the seven columns the policy names, with data under them")
def required_columns(workspace: Path) -> float:
    sheets = _populated_sheets(workspace, "required_columns")
    if sheets is None:
        return 0.0
    best = 0.0
    for _name, header in sheets:
        present = {str(c).strip().lower() for c in header if c is not None}
        best = max(best, sum(1 for c in COLUMNS if c in present) / len(COLUMNS))
    if best < 1.0:
        print(f"required_columns: best populated sheet matched {best:.2f} of the required header")
    return best


@criterion(description="the column order section 7.1 fixes")
def column_order(workspace: Path) -> float:
    sheets = _populated_sheets(workspace, "column_order")
    if sheets is None:
        return 0.0
    best, worst = 0.0, sheets[0][0]
    for name, header in sheets:
        normalised = [str(c).strip().lower() if c is not None else "" for c in header]
        hits = sum(1 for i, col in enumerate(COLUMNS)
                   if i < len(normalised) and normalised[i] == col)
        if hits / len(COLUMNS) > best:
            best = hits / len(COLUMNS)
            worst = name
    if best < 1.0:
        print(f"column_order: best populated sheet ({worst!r}) placed {best:.2f} of the seven "
              f"columns at the position section 7.1 fixes: {list(COLUMNS)}")
    return best


@criterion(description="the sheet name section 7.1 fixes")
def sheet_name(workspace: Path) -> float:
    sheets = _populated_sheets(workspace, "sheet_name")
    if sheets is None:
        return 0.0
    for name, header in sheets:
        present = {str(c).strip().lower() for c in header if c is not None}
        if name.strip() == SHEET and all(c in present for c in COLUMNS):
            return 1.0
    print(f"sheet_name: no populated sheet carrying the required header is named {SHEET!r}; "
          f"found {[n for n, _ in sheets]}")
    return 0.0


@criterion(description="workings.json parses and carries the keys the policy names, against a real restatement")
def required_keys(workspace: Path) -> float:
    import json

    path = workspace / "workings.json"
    if not path.exists():
        print("required_keys: workings.json was not produced")
        return 0.0
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        print(f"required_keys: workings.json is not valid JSON: {type(exc).__name__}: {exc}")
        return 0.0
    if not isinstance(payload, dict):
        print(f"required_keys: workings.json holds {type(payload).__name__}, not an object")
        return 0.0

    # the data-row gate, in JSON: transactions_by_currency carries one key per currency
    # in the population, so an empty map is the shell rather than the restatement.
    counts = payload.get("transactions_by_currency")
    if not isinstance(counts, dict) or not counts:
        print("required_keys: transactions_by_currency is empty, so workings.json is a shell "
              "rather than a restatement; section 7.2 asks for one key per currency in the "
              "population")
        return 0.0
    return sum(1 for k in KEYS if k in payload) / len(KEYS)
