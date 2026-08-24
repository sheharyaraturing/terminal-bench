"""The two promised artifacts appeared, in the layout SOP sections 9 and 11 state.

Structure only: this dimension separates "did nothing" from "did something wrong".
It never inspects a computed value, so it cannot pay for a plausible-looking report.

It also never pays for an empty shell. Each criterion requires at least one data row
below the header — SOP section 11 asks for one row per in-scope ticket and section 9
for one row per ticket meeting a condition, so a file carrying only headers has not
satisfied either, and a headers-only workbook scores 0 here rather than banking a
quarter of the reward for free.

The layout requirements the SOP states are graded rather than assumed: the sheet is
named SLA, and the header names sit in the stated order. Both are folded into one
fraction so a partly-right header reads as partly right.
"""

import csv
from pathlib import Path

from rewardkit import criterion

REPORT_SHEET = "SLA"
REPORT_COLUMNS = ("ticket", "severity", "deadline", "breached")
LOG_COLUMNS = ("ticket", "reason")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _positional(header: list[str], columns: tuple[str, ...]) -> int:
    """How many required column names sit at the index the SOP puts them at."""
    return sum(1 for i, col in enumerate(columns)
               if i < len(header) and header[i].lower() == col)


@criterion(description="sla_report.xlsx has a sheet named SLA headed Ticket, Severity, "
                       "Deadline, Breached in that order, over at least one data row")
def report_layout(workspace: Path) -> float:
    import openpyxl

    path = workspace / "sla_report.xlsx"
    if not path.exists():
        print("report_layout: sla_report.xlsx was never written")
        return 0.0
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"report_layout: could not open {path.name}: {type(exc).__name__}: {exc}")
        return 0.0

    best = 0.0
    header_only: list[str] = []
    # never trust .active — an empty leading sheet is a legal thing for an agent
    # to leave behind, and the graded sheet may be any of them
    for name in wb.sheetnames:
        rows = [r for r in wb[name].iter_rows(values_only=True)]
        if not rows:
            continue
        header = [_text(c) for c in rows[0]]
        if not _positional(header, REPORT_COLUMNS) and not any(
            _text(c).lower() in REPORT_COLUMNS for c in rows[0]
        ):
            continue
        data = [r for r in rows[1:] if r and any(_text(c) for c in r)]
        if not data:
            header_only.append(name)
            continue
        hit = _positional(header, REPORT_COLUMNS)
        named = 1 if name.strip().upper() == REPORT_SHEET else 0
        best = max(best, (hit + named) / (len(REPORT_COLUMNS) + 1))

    if header_only and best == 0.0:
        print(f"report_layout: sheet(s) {header_only} carry the header and no data row; "
              "SOP section 11 asks for one row per in-scope ticket")
    elif best < 1.0:
        print(f"report_layout: no sheet named {REPORT_SHEET} in {path.name} heads its "
              f"data with {REPORT_COLUMNS} in that order; sheets are {wb.sheetnames}")
    return best


@criterion(description="sla_exceptions.csv is headed Ticket,Reason in that order, "
                       "over at least one data row")
def log_layout(workspace: Path) -> float:
    path = workspace / "sla_exceptions.csv"
    if not path.exists():
        print("log_layout: sla_exceptions.csv was never written")
        return 0.0
    try:
        with path.open(newline="", encoding="utf-8-sig", errors="replace") as fh:
            rows = [r for r in csv.reader(fh) if any(_text(c) for c in r)]
    except Exception as exc:
        print(f"log_layout: could not read {path.name}: {type(exc).__name__}: {exc}")
        return 0.0

    for i, row in enumerate(rows[:3]):
        header = [_text(c) for c in row]
        hit = _positional(header, LOG_COLUMNS)
        if not hit and not any(c.lower() in LOG_COLUMNS for c in header):
            continue
        if not rows[i + 1:]:
            print(f"log_layout: {path.name} carries the header and no data row; SOP "
                  "section 9 asks for one row per in-scope ticket meeting a condition")
            return 0.0
        if hit < len(LOG_COLUMNS):
            print(f"log_layout: header {row} is not {LOG_COLUMNS} in that order")
        return hit / len(LOG_COLUMNS)

    print(f"log_layout: {path.name} carries no Ticket,Reason header; first row was "
          f"{rows[0] if rows else '(empty file)'}")
    return 0.0
