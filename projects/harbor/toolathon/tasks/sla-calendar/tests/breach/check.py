"""Fraction of tickets whose breach flag matches the SOP's rule.

Breach is graded apart from the deadline because the two fail for different
reasons: a wrong deadline is arithmetic, a wrong flag with a right deadline is a
misread of which instant governs — the first response where there is one, the
pinned as-of instant where there is not.
"""

from datetime import datetime
from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/sla_report.xlsx")
REQUIRED = ("ticket", "severity", "deadline", "breached")
FORMULA = "<spreadsheet formula, no stored value>"


def _cell(value) -> str:
    """One agent-written cell as comparable text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "YES" if value else "NO"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read(path: Path, label: str) -> dict[str, dict[str, str]]:
    """Rows keyed by ticket, taken from whichever sheet carries the header.

    The workbook is loaded twice. `data_only=True` returns cached formula results,
    and a workbook written by a program that never ran Excel has no cache, so every
    formula cell would otherwise read blank. The second, formula-bearing load lets
    a formula be reported as a formula rather than misdiagnosed as an empty cell,
    and the first such cell is named by sheet and coordinate so whoever debugs the
    run is not sent after a phantom blank.
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    if not path.exists():
        print(f"{label}: {path.name} does not exist")
        return {}
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"{label}: could not open {path.name}: {type(exc).__name__}: {exc}")
        return {}
    try:
        literal = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        literal = None

    for name in cached.sheetnames:
        rows = [list(r) for r in cached[name].iter_rows(values_only=True)]
        if not rows:
            continue
        header = [_cell(c).lower() for c in rows[0]]
        if any(col not in header for col in REQUIRED):
            continue
        raw = []
        if literal is not None and name in literal.sheetnames:
            raw = [list(r) for r in literal[name].iter_rows(values_only=True)]

        idx = {col: header.index(col) for col in REQUIRED}
        out: dict[str, dict[str, str]] = {}
        formulas = 0
        first_formula: tuple[str, str] | None = None
        for i, row in enumerate(rows[1:], start=1):
            if not row or all(c is None for c in row):
                continue
            record = {}
            for col, j in idx.items():
                value = row[j] if j < len(row) else None
                if value is None and i < len(raw) and j < len(raw[i]):
                    alt = raw[i][j]
                    if isinstance(alt, str) and alt.startswith("="):
                        if first_formula is None:
                            # iter_rows starts at A1, so data row i is sheet row i+1
                            first_formula = (f"{get_column_letter(j + 1)}{i + 1}", alt)
                        value, formulas = FORMULA, formulas + 1
                    elif alt is not None:
                        value = alt
                record[col] = _cell(value)
            ticket = record["ticket"].upper()
            if ticket:
                out[ticket] = record
        if first_formula is not None:
            # one line per affected sheet, not per cell: a wholly formula-written
            # answer should name its cause once rather than flood the log
            coord, expr = first_formula
            print(f"{label}: {name}!{coord} holds the formula {expr!r} and no cached "
                  f"value; SOP section 11 requires recorded values, not formulas "
                  f"({formulas} such cell(s) in this sheet)")
        return out

    print(f"{label}: no sheet in {path.name} carries all of {REQUIRED}; "
          f"sheets are {cached.sheetnames}")
    return {}


def _pair(workspace: Path):
    return _read(EXPECTED, "expected"), _read(workspace / "sla_report.xlsx", "report")


@criterion(description="each ticket is flagged breached or not as the SOP defines it")
def breached(workspace: Path) -> float:
    want, got = _pair(workspace)
    if not want:
        return 0.0
    hits = 0
    misses: list[str] = []
    for ticket, row in want.items():
        mine = got.get(ticket, {}).get("breached", "")
        if mine.upper() == row["breached"].upper():
            hits += 1
        elif len(misses) < 6:
            misses.append(f"{ticket}: wanted {row['breached']!r}, got {mine!r}")
    if misses:
        print("breached: " + "; ".join(misses))
    return hits / max(len(want), len(got))
