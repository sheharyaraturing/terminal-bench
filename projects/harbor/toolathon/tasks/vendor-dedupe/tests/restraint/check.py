"""The cluster the standard refuses to resolve is flagged rather than guessed.

One cluster in the register carries two irreconcilable tax identifiers. The standard
says to leave its records alone and report them; an agent that picks one identifier,
averages the rows, or quietly drops the cluster from the log has guessed.

Which cluster that is comes from the answer key rather than from this file: it is the
one whose members all still stand as separate rows in the consolidated register.

Four things are counted, and each must be present *and* right -- a missing log row
scores nothing, as do the near-miss spellings of "nothing was taken here": an empty
cell, a zero, a null, and an omitted column are all rejected, because the standard
publishes two distinct markers and only one of them belongs on a held cluster.
"""

import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from rewardkit import criterion

EXPECTED_LOG = Path("/tests/expected/merge_log.xlsx")
EXPECTED_REGISTER = Path("/tests/expected/vendor_master_merged.xlsx")
LOG_COLUMNS = ("Survivor", "Losers", "FieldsTaken", "Status")
REGISTER_COLUMNS = (
    "VendorID", "LegalName", "TaxID", "RemitToAddress", "RemitToPostcode",
    "Country", "PaymentTerms", "ContactEmail", "DeptCode", "SpendYTD", "LastUpdated",
)


def _scalar(value) -> str:
    """One cell, reduced to the form both sides are compared in."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, (int, float, Decimal)):
        text = format(Decimal(str(value)), "f")
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        return text
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return " ".join(str(value).split()).casefold()


def _table(path: Path, required: tuple[str, ...], label: str):
    """Header and rows of the first sheet carrying every required column.

    Read twice. `data_only=True` returns Excel's cached values, and a workbook
    written by a program that never ran Excel has no cache, so a formula cell reads
    as None; the second read supplies whatever the cell actually holds, so a formula
    cell arrives here as its formula text rather than as an empty one.

    Formulas are not evaluated. openpyxl cannot evaluate them and this verifier has
    no formula engine, and section 9 of the standard asks for recorded values, so a
    formula cell is simply compared as the text it holds and scores wrong. What is
    reported is the cause: the first offending cell on each sheet is named by its
    coordinate, once per sheet rather than once per cell, so a wholly formula-written
    answer names the problem instead of flooding the log with phantom blanks.
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    if not path.exists():
        print(f"{label}: {path.name} was not produced")
        return None, []
    try:
        cached = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"{label}: could not open {path.name}: {type(exc).__name__}: {exc}")
        return None, []
    try:
        raw = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        raw = cached

    wanted = {c.casefold() for c in required}
    for name in cached.sheetnames:
        rows_c = list(cached[name].iter_rows(values_only=True))
        rows_r = list(raw[name].iter_rows(values_only=True)) if name in raw.sheetnames else rows_c
        if not rows_c:
            continue
        header = ["".join(str(c).split()).casefold() if c is not None else "" for c in rows_c[0]]
        if not wanted.issubset(set(header)):
            continue
        index = {col: header.index(col.casefold()) for col in required if col.casefold() in header}
        out = []
        formulas: list[tuple[tuple[int, int], str, str]] = []
        for i, row_c in enumerate(rows_c[1:], start=1):
            row_r = rows_r[i] if i < len(rows_r) else row_c
            if row_c is None or all(c is None for c in row_c):
                continue
            record = {}
            for col, pos in index.items():
                value = row_c[pos] if pos < len(row_c) else None
                if value is None and pos < len(row_r) and row_r[pos] is not None:
                    value = row_r[pos]
                    if isinstance(value, str) and value.startswith("="):
                        coord = f"{get_column_letter(pos + 1)}{i + 1}"
                        formulas.append(((i, pos), coord, value))
                record[col] = value
            out.append(record)
        if formulas:
            formulas.sort(key=lambda f: f[0])
            _, coord, text = formulas[0]
            print(f"{label}: {name}!{coord} holds the formula {text!r} and no cached "
                  "value; section 9 of Data Governance Standard DG-114 requires "
                  "recorded values, not formulas")
            if len(formulas) > 1:
                print(f"{label}: {len(formulas) - 1} further formula cell(s) on sheet "
                      f"{name!r} are read the same way and are not reported "
                      "individually")
        return name, out
    print(f"{label}: no sheet in {path.name} carries the columns {list(required)}")
    return None, []


def _list(value):
    text = _scalar(value)
    if not text:
        return None
    parts = [p.strip() for p in re.split(r"[;,]", text) if p.strip()]
    return tuple(sorted(parts)) if parts else None


def _log(path: Path, label: str) -> dict[str, dict]:
    _, rows = _table(path, LOG_COLUMNS, label)
    return {_scalar(r.get("Survivor")): r for r in rows if _scalar(r.get("Survivor"))}


def _register(path: Path, label: str) -> dict[str, tuple[str, ...]]:
    _, rows = _table(path, REGISTER_COLUMNS, label)
    out: dict[str, tuple[str, ...]] = {}
    for row in rows:
        vendor = _scalar(row.get("VendorID"))
        if vendor:
            out[vendor] = tuple(_scalar(row.get(c)) for c in REGISTER_COLUMNS)
    return out


@criterion(description="the cluster with irreconcilable tax identifiers is flagged, not resolved")
def restraint(workspace: Path) -> float:
    want_log = _log(EXPECTED_LOG, "restraint/expected")
    want_register = _register(EXPECTED_REGISTER, "restraint/expected")
    if not want_log or not want_register:
        return 0.0

    held = []
    for survivor, row in want_log.items():
        members = [survivor, *(_list(row.get("Losers")) or ())]
        if all(m in want_register for m in members):
            held.append((survivor, row, members))
    if not held:
        print("restraint: the answer key holds no un-consolidated cluster to grade")
        return 0.0

    got_log = _log(workspace / "merge_log.xlsx", "restraint")
    got_register = _register(workspace / "vendor_master_merged.xlsx", "restraint")

    hits = 0
    for survivor, want_row, members in held:
        have = got_log.get(survivor)
        if have is None:
            print(f"restraint: {survivor} has no merge-log row; the cluster was not reported")
        else:
            if _scalar(have.get("Status")) == _scalar(want_row.get("Status")):
                hits += 1
            else:
                print(f"restraint: {survivor} status {_scalar(have.get('Status'))!r}, "
                      f"expected {_scalar(want_row.get('Status'))!r}")
            marker = _list(want_row.get("FieldsTaken"))
            if _list(have.get("FieldsTaken")) == marker:
                hits += 1
            else:
                print(f"restraint: {survivor} FieldsTaken {have.get('FieldsTaken')!r} is not "
                      f"the published marker {want_row.get('FieldsTaken')!r}")
            if _list(have.get("Losers")) == _list(want_row.get("Losers")):
                hits += 1
            else:
                print(f"restraint: {survivor} losers {have.get('Losers')!r}, "
                      f"expected {want_row.get('Losers')!r}")
        kept = [m for m in members if got_register.get(m) == want_register[m]]
        if len(kept) == len(members) and have is not None:
            hits += 1
        elif len(kept) != len(members):
            missing = [m for m in members if m not in kept]
            print(f"restraint: {missing} no longer stand as their own rows in the "
                  "consolidated register")
        else:
            print(f"restraint: {members} were left alone but never reported; "
                  "leaving a conflict unexamined is not the same as flagging it")
    return hits / (4 * len(held))
