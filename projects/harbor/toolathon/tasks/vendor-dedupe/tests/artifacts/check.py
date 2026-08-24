"""Both deliverables appeared, carry the columns the instruction names, and hold data.

Sheets are found by scanning every worksheet for the required header, never by
trusting `.active`: an agent that leaves an empty `Sheet1` in front of its answer
has still produced the answer.

A header row on its own is not an answer. A sheet counts here only when at least one
data row sits under its header, so the empty shell -- correct column names, nothing
beneath them -- scores nothing rather than banking a dimension for free. File
existence is not scored on its own for the same reason: a file that exists and says
nothing is worth what it says.
"""

from pathlib import Path

from rewardkit import criterion

REGISTER_COLUMNS = (
    "vendorid", "legalname", "taxid", "remittoaddress", "remittopostcode",
    "country", "paymentterms", "contactemail", "deptcode", "spendytd", "lastupdated",
)
LOG_COLUMNS = ("survivor", "losers", "fieldstaken", "status")


def _populated_headers(path: Path) -> set[str]:
    """Header cells of every sheet that carries at least one data row, normalised."""
    import openpyxl

    if not path.exists():
        print(f"artifacts: {path.name} was not produced")
        return set()
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"artifacts: could not open {path.name}: {type(exc).__name__}: {exc}")
        return set()
    found: set[str] = set()
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        if not rows:
            continue
        header = {"".join(str(c).split()).lower() for c in rows[0] if c is not None}
        if not header:
            continue
        data = [r for r in rows[1:]
                if r is not None and any(c is not None and str(c).strip() for c in r)]
        if not data:
            print(f"artifacts: {path.name} sheet {name!r} carries a header row and no "
                  "data rows; a header row on its own records no answer")
            continue
        found |= header
    return found


def _fraction(path: Path, required: tuple[str, ...]) -> float:
    found = _populated_headers(path)
    missing = [c for c in required if c not in found]
    if missing:
        print(f"artifacts: {path.name} is missing header(s) {missing} from any sheet "
              "that carries data")
    return (len(required) - len(missing)) / len(required)


@criterion(description="vendor_master_merged.xlsx carries the register's columns over data rows")
def merged_register_columns(workspace: Path) -> float:
    return _fraction(workspace / "vendor_master_merged.xlsx", REGISTER_COLUMNS)


@criterion(description="merge_log.xlsx carries Survivor, Losers, FieldsTaken and Status over data rows")
def merge_log_columns(workspace: Path) -> float:
    return _fraction(workspace / "merge_log.xlsx", LOG_COLUMNS)
