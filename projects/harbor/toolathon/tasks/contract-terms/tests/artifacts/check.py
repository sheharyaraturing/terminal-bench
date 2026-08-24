"""The schedule exists, is laid out as clause 7 requires, and names its suppliers.

Three criteria, none of which a headers-only shell can bank:

* `schedule_shape` wants the worksheet clause 7.1 names, the four columns in the
  order clause 7.2 fixes, and at least one row of data beneath them. A workbook
  carrying nothing but the header row answers no invoice, so it scores zero here
  as it does on the two content dimensions.
* `row_order` grades clause 7.3's ascending-Invoice requirement position by
  position. The suite's rule is that a stated ordering is compared positionally
  or is not stated at all; the manual states it, so it is measured.
* `vendor_labels` grades the trading name clause 7.4 fixes, so a sheet carrying
  the header and junk beneath it is not a schedule.

`.active` is not trusted for the content scan: an agent that leaves an empty
`Sheet1` in front of its work is not wrong, so every sheet is searched for the
header row.
"""

from pathlib import Path

from rewardkit import criterion

EXPECTED = Path("/tests/expected/payment_schedule.xlsx")
ARTIFACT = "payment_schedule.xlsx"
REQUIRED = ("invoice", "vendor", "duedate", "governingsource")
SHEET = "schedule"


def _norm_head(value) -> str:
    return "" if value is None else str(value).strip().lower().replace(" ", "").replace("_", "")


def _unwrap_formula(text: str) -> str:
    """A formula cell that is really a literal, e.g. ="Corvid Logistics"."""
    body = text.lstrip("=").strip()
    if len(body) >= 2 and body[0] == body[-1] and body[0] in "\"'":
        return body[1:-1]
    return text


def _open(path: Path, label: str):
    import openpyxl

    if not path.exists():
        print(f"artifacts: {label} {path.name} is missing")
        return None, None
    try:
        values = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        print(f"artifacts: could not open {label} {path.name}: {type(exc).__name__}: {exc}")
        return None, None
    try:
        formulas = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        formulas = None
    return values, formulas


def _vendors(path: Path, label: str) -> tuple[dict[str, str], int, list[str]]:
    """Trading name per invoice, the number of data rows, and the order written.

    Scans every sheet for the header rather than trusting `.active`, and keeps a
    formula-bearing read so a `="Corvid Logistics"` cell is reported as the
    formula it is rather than read back as blank.
    """
    values, formulas = _open(path, label)
    if values is None:
        return {}, 0, []

    for sheet in values.worksheets:
        head_row = None
        columns: dict[str, int] = {}
        for row in sheet.iter_rows(min_row=1, max_row=10):
            names: dict[str, int] = {}
            for cell in row:
                name = _norm_head(cell.value)
                if name in REQUIRED and name not in names:
                    names[name] = cell.column
            if all(c in names for c in REQUIRED):
                head_row, columns = row[0].row, names
                break
        if head_row is None:
            continue

        mirror = None
        if formulas is not None and sheet.title in formulas.sheetnames:
            mirror = formulas[sheet.title]

        out: dict[str, str] = {}
        order: list[str] = []
        rows = 0
        held: list[tuple[str, str]] = []
        for r in range(head_row + 1, sheet.max_row + 1):
            raw_key = sheet.cell(row=r, column=columns["invoice"]).value
            cell = sheet.cell(row=r, column=columns["vendor"])
            raw_vendor = cell.value
            if raw_key is None and raw_vendor is None:
                continue
            if raw_vendor is None and mirror is not None:
                shadow = mirror.cell(row=r, column=columns["vendor"]).value
                if isinstance(shadow, str) and shadow.startswith("="):
                    held.append((cell.coordinate, shadow))
                    raw_vendor = _unwrap_formula(shadow)
            rows += 1
            key = "" if raw_key is None else str(raw_key).strip().upper()
            if key:
                order.append(key)
                out[key] = "" if raw_vendor is None else " ".join(str(raw_vendor).split()).upper()
        if held:
            coord, formula = held[0]
            extra = f" ({len(held)} such cells on this sheet)" if len(held) > 1 else ""
            print(
                f"artifacts: {label} {sheet.title}!{coord} holds the formula {formula!r} and no "
                f"cached value; AP manual clause 7.6 requires recorded values, not formulas{extra}"
            )
        return out, rows, order

    print(f"artifacts: no sheet in {label} {path.name} carries all of {REQUIRED}")
    return {}, 0, []


@criterion(
    description="payment_schedule.xlsx has a Schedule sheet, the four columns in order and data"
)
def schedule_shape(workspace: Path) -> float:
    path = workspace / ARTIFACT
    values, _ = _open(path, "agent")
    if values is None:
        print(f"artifacts: {ARTIFACT} was not produced")
        return 0.0

    sheet = None
    for candidate in values.worksheets:
        if candidate.title.strip().lower() == SHEET:
            sheet = candidate
            break
    if sheet is None:
        print(
            f"artifacts: no worksheet named Schedule; found {values.sheetnames} — "
            f"AP manual clause 7.1 names the sheet"
        )
        return 0.0

    header = [_norm_head(c) for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
    hits = sum(1 for i, name in enumerate(REQUIRED) if i < len(header) and header[i] == name)
    if hits < len(REQUIRED):
        print(
            f"artifacts: row 1 of Schedule reads {header[:6]}; AP manual clause 7.2 fixes "
            f"{list(REQUIRED)} in that order"
        )

    written = 0
    for row in sheet.iter_rows(min_row=2, max_col=max(4, len(header)), values_only=True):
        if any(c is not None and str(c).strip() != "" for c in row):
            written += 1
    if written == 0:
        print(
            "artifacts: Schedule carries the header row and no data row; every in-scope invoice "
            "is unanswered (AP manual clause 6.3)"
        )
        return 0.0
    return hits / len(REQUIRED)


@criterion(description="scheduled rows sit in ascending order of Invoice")
def row_order(workspace: Path) -> float:
    _, _, expected = _vendors(EXPECTED, "expected")
    if not expected:
        return 0.0
    _, rows, got = _vendors(workspace / ARTIFACT, "agent")
    hits = sum(1 for i, ref in enumerate(expected) if i < len(got) and got[i] == ref)
    if hits < len(expected):
        first = next(
            (i for i, ref in enumerate(expected) if i >= len(got) or got[i] != ref), None
        )
        seen = got[first] if first is not None and first < len(got) else "(nothing)"
        print(
            f"artifacts: {len(expected) - hits} rows out of position; position {first + 1} holds "
            f"{seen} where {expected[first]} belongs — AP manual clause 7.3 requires ascending "
            f"order of Invoice"
        )
    return hits / max(len(expected), rows, len(got), 1)


@criterion(description="each scheduled row names its supplier's trading name")
def vendor_labels(workspace: Path) -> float:
    expected, _, _ = _vendors(EXPECTED, "expected")
    if not expected:
        return 0.0
    got, rows, _ = _vendors(workspace / ARTIFACT, "agent")
    hits = sum(1 for ref, name in expected.items() if got.get(ref, "") == name)
    misses = [ref for ref, name in expected.items() if got.get(ref, "") != name]
    if misses:
        print(f"artifacts: vendor wrong or missing on {len(misses)}, first few {misses[:5]}")
    return hits / max(len(expected), rows, 1)
